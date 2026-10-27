"""Importa requisitos/subitens de uma norma a partir de uma lista de verificação em Excel.

Cada coluna do cabeçalho recebe um papel (`PAPEIS`): código do item, pergunta/título, orientação,
classificação (sigla, ex.: M/I/D), área por marcação (a coluna é a área; célula preenchida = a pergunta
pertence a ela, ex.: SASSMAQ SS/MA/Q), área pelo valor da célula, ou ignorar. O papel é sugerido pelo
mapeamento salvo na norma, pelo nome do cabeçalho ou pelas siglas de área da norma, e pode ser trocado
na análise; o mapeamento usado na importação fica salvo na norma.

A hierarquia vem do código: ``1`` > ``1.1`` > ``1.1.1`` > ``1.1.1.a`` (``1.2.1a`` também é filho de
``1.2.1``). Itens com filhos viram texto orientativo; itens sem filhos viram perguntas.
"""

import re
import unicodedata
from io import BytesIO

import openpyxl
from django.db import transaction
from openpyxl.styles import Alignment, Font, PatternFill
from django.db.models import Max

from .models import ItemNormaSGQ, RespostaAuditoriaSGQ

_SUBSTITUIR_RESPONDIDA = (
    'A norma tem perguntas já respondidas em auditorias; os itens atuais não podem ser substituídos, só complementados.'
)

PAPEIS = ('codigo', 'titulo', 'descricao', 'classificacao', 'area', 'areaValor', 'ignorar')
_PAPEIS_UNICOS = ('codigo', 'titulo', 'descricao', 'classificacao', 'areaValor')
_NOMES_PAPEIS = {
    'codigo': 'código do item',
    'titulo': 'pergunta/título',
    'descricao': 'orientação',
    'classificacao': 'classificação',
    'areaValor': 'área (valor da célula)',
}

_CODIGO_RE = re.compile(r'^\d+(?:[.\s]*(?:\d+|[A-Za-z]))*\.?$')
_TOKEN_RE = re.compile(r'\d+|[A-Za-z]+')
_SIGLA_RE = re.compile(r'^[A-Z0-9]+')
_HEADERS_CODIGO = ('requisito', 'codigo', 'item', 'numero', 'n')
_HEADERS_TITULO = ('lista de verificacao', 'titulo', 'pergunta', 'questao', 'verificacao')
_HEADERS_DESCRICAO = ('guia', 'orientac', 'descricao', 'observac')
_HEADERS_CLASSIFICACAO = ('tipo', 'classificac')
_HEADERS_AREA = ('area', 'submodulo', 'modulo')
_NAO_MARCADO = ('', '0', 'nao', 'n', 'false', '-')
_LINHAS_CABECALHO = 60
_MAX_TITULO = ItemNormaSGQ._meta.get_field('titulo').max_length
_MAX_DESCRICAO = 4000
_MAX_SIGLA = 10


class NormaItensImportError(Exception):
    pass


def _normalizar(valor) -> str:
    texto = unicodedata.normalize('NFKD', str(valor or '')).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', texto).strip().lower()


def _texto(valor) -> str:
    return re.sub(r'[ \t]+', ' ', str(valor or '')).strip()


def _sigla(valor) -> str:
    """Sigla no início da célula: "M", "M - Mandatória" e "m" viram "M"."""
    encontrada = _SIGLA_RE.match(_normalizar(valor).upper())
    return encontrada.group(0)[:_MAX_SIGLA] if encontrada else ''


def _celula(row, coluna):
    return row[coluna] if coluna is not None and coluna < len(row) else None


def _ler_linhas(file_bytes: bytes) -> list[tuple]:
    try:
        workbook = openpyxl.load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise NormaItensImportError('Não foi possível abrir o arquivo. Envie uma planilha Excel (.xlsx).') from exc
    rows = list(workbook.worksheets[0].iter_rows(values_only=True))
    workbook.close()
    return rows


def _papeis_por_nome(cabecalhos: list[str]) -> dict[int, str]:
    papeis: dict[int, str] = {}

    def primeira(condicao):
        return next((i for i, h in enumerate(cabecalhos) if h and i not in papeis and condicao(h)), None)

    for papel, condicao in (
        ('codigo', lambda h: h in _HEADERS_CODIGO),
        ('titulo', lambda h: h.startswith(_HEADERS_TITULO)),
        ('descricao', lambda h: h.startswith(_HEADERS_DESCRICAO)),
        ('classificacao', lambda h: h.startswith(_HEADERS_CLASSIFICACAO)),
        ('areaValor', lambda h: h.startswith(_HEADERS_AREA)),
    ):
        coluna = primeira(condicao)
        if coluna is not None:
            papeis[coluna] = papel
    return papeis


def _detectar_cabecalho(rows, salvo: list[dict]) -> int:
    salvos = {item['papel']: item['cabecalho'] for item in salvo if item.get('papel') in ('codigo', 'titulo')}
    for indice, row in enumerate(rows[:_LINHAS_CABECALHO]):
        cabecalhos = [_normalizar(valor) for valor in row]
        papeis = set(_papeis_por_nome(cabecalhos).values())
        if {'codigo', 'titulo'} <= papeis:
            return indice
        if len(salvos) == 2 and set(salvos.values()) <= set(cabecalhos):
            return indice
    for indice, row in enumerate(rows[:_LINHAS_CABECALHO]):
        if sum(1 for valor in row if _texto(valor)) >= 2:
            return indice
    raise NormaItensImportError('Cabeçalho não encontrado. A planilha precisa de uma linha com os nomes das colunas.')


def _exemplos(rows, linha: int, coluna: int, cabecalho: str) -> list[str]:
    vistos: list[str] = []
    for row in rows[linha + 1:]:
        valor = _texto(_celula(row, coluna))
        if valor and _normalizar(valor) != cabecalho and valor not in vistos:
            vistos.append(valor[:40])
            if len(vistos) == 3:
                break
    return vistos


def _colunas(rows, linha: int, norma, mapeamento: list[dict] | None) -> list[dict]:
    cabecalho = rows[linha]
    normalizados = [_normalizar(valor) for valor in cabecalho]
    salvo = {item['cabecalho']: item for item in norma.mapeamento_importacao}
    por_nome = _papeis_por_nome(normalizados)
    siglas_area = norma.siglas_areas()
    escolhidos = {int(item['indice']): item for item in mapeamento or [] if str(item.get('indice', '')).isdigit()}

    colunas = []
    for indice, valor in enumerate(cabecalho):
        chave = normalizados[indice]
        if not chave:
            continue
        sigla = _sigla(valor)
        if indice in escolhidos:
            papel, area = escolhidos[indice].get('papel'), _sigla(escolhidos[indice].get('area') or sigla)
        elif chave in salvo:
            papel, area = salvo[chave].get('papel'), salvo[chave].get('area') or sigla
        elif indice in por_nome:
            papel, area = por_nome[indice], sigla
        elif sigla and sigla in siglas_area and chave.upper() == sigla:
            papel, area = 'area', sigla
        else:
            papel, area = 'ignorar', sigla
        if papel not in PAPEIS:
            raise NormaItensImportError(f'Papel inválido na coluna "{_texto(valor)}".')
        if papel == 'area' and not area:
            raise NormaItensImportError(f'Informe a sigla da área da coluna "{_texto(valor)}".')
        colunas.append({
            'indice': indice,
            'cabecalho': _texto(valor),
            'papel': papel,
            'area': area,
            'exemplos': _exemplos(rows, linha, indice, chave),
        })

    for papel in _PAPEIS_UNICOS:
        usados = [coluna['cabecalho'] for coluna in colunas if coluna['papel'] == papel]
        if len(usados) > 1:
            raise NormaItensImportError(
                f'Só uma coluna pode ser {_NOMES_PAPEIS[papel]} (marcadas: {", ".join(usados)}).'
            )
    return colunas


def _montar_nos(rows, linha: int, colunas: list[dict]) -> dict:
    por_papel = {coluna['papel']: coluna for coluna in colunas if coluna['papel'] in _PAPEIS_UNICOS}
    colunas_area = [coluna for coluna in colunas if coluna['papel'] == 'area']

    def valor(row, papel):
        coluna = por_papel.get(papel)
        if coluna is None:
            return None
        bruto = _celula(row, coluna['indice'])
        # Listas como a do SASSMAQ repetem o cabeçalho ("Tipo", "SS") nas linhas de seção.
        return None if _normalizar(bruto) == _normalizar(coluna['cabecalho']) else bruto

    def areas(row) -> list[str]:
        marcadas = []
        for coluna in colunas_area:
            celula = _normalizar(_celula(row, coluna['indice']))
            if celula not in _NAO_MARCADO and celula != _normalizar(coluna['cabecalho']):
                marcadas.append(coluna['area'])
        da_celula = _texto(valor(row, 'areaValor')).upper()[:_MAX_SIGLA]
        if da_celula:
            marcadas.append(da_celula)
        return list(dict.fromkeys(marcadas))

    nos: list[dict] = []
    por_codigo: dict[tuple, int] = {}
    ignoradas: list[int] = []
    avisos: list[str] = []

    for numero_linha, row in enumerate(rows[linha + 1:], start=linha + 2):
        codigo = _texto(valor(row, 'codigo'))
        titulo = _texto(valor(row, 'titulo'))
        if not codigo and not titulo and not any(v not in (None, '') for v in row):
            continue
        if not _CODIGO_RE.match(codigo) or not titulo:
            ignoradas.append(numero_linha)
            continue

        tokens = tuple(t.lower() for t in _TOKEN_RE.findall(codigo))
        parent = None
        for tamanho in range(len(tokens) - 1, 0, -1):
            if tokens[:tamanho] in por_codigo:
                parent = por_codigo[tokens[:tamanho]]
                break
        if parent is None and len(tokens) > 1:
            avisos.append(f'Linha {numero_linha}: item pai de {codigo} não está na planilha; importado como requisito principal.')

        if len(titulo) > _MAX_TITULO:
            avisos.append(f'Linha {numero_linha} ({codigo}): título com mais de {_MAX_TITULO} caracteres foi cortado.')
            titulo = titulo[:_MAX_TITULO]
        descricao = _texto(valor(row, 'descricao'))
        if descricao == titulo:
            descricao = ''
        if len(descricao) > _MAX_DESCRICAO:
            avisos.append(f'Linha {numero_linha} ({codigo}): guia com mais de {_MAX_DESCRICAO} caracteres foi cortado.')
            descricao = descricao[:_MAX_DESCRICAO]

        if tokens in por_codigo:
            avisos.append(f'Linha {numero_linha}: código {codigo} repetido na planilha.')
        else:
            por_codigo[tokens] = len(nos)
        nos.append({
            'codigo': codigo.rstrip('.'),
            'titulo': titulo,
            'descricao': descricao,
            'classificacao': _sigla(valor(row, 'classificacao')),
            'areas': areas(row),
            'parent': parent,
            'filhos': 0,
        })
        if parent is not None:
            nos[parent]['filhos'] += 1

    if not nos:
        raise NormaItensImportError('Nenhum item com código encontrado na planilha. Confira a coluna de código do item.')
    return {'nos': nos, 'ignoradas': ignoradas, 'avisos': avisos}


def _contar(valores) -> dict[str, int]:
    contagem: dict[str, int] = {}
    for valor in valores:
        contagem[valor] = contagem.get(valor, 0) + 1
    return contagem


def importar_itens_norma(
    norma,
    file_bytes: bytes,
    *,
    tipo_resposta,
    substituir: bool,
    dry_run: bool,
    mapeamento: list[dict] | None = None,
    novas_classificacoes: list[dict] | None = None,
    novas_areas: list[dict] | None = None,
) -> dict:
    """`tipo_resposta` é o TipoRespostaSGQ (da mesma filial da norma) aplicado a todas as perguntas.

    A análise devolve as colunas com o papel de cada uma e as classificações/áreas da planilha que a norma
    ainda não tem; na importação elas precisam vir definidas e são gravadas na norma junto com o mapeamento.
    """
    rows = _ler_linhas(file_bytes)
    linha = _detectar_cabecalho(rows, norma.mapeamento_importacao)
    colunas = _colunas(rows, linha, norma, mapeamento)
    resumo = {
        'linhaCabecalho': linha + 1,
        'colunas': colunas,
        'pendencia': '',
        'total': 0,
        'requisitos': 0,
        'orientativos': 0,
        'perguntas': 0,
        'porClassificacao': {},
        'classificacoesNovas': [],
        'porArea': {},
        'areasNovas': [],
        'linhasIgnoradas': [],
        'avisos': [],
        'itensExistentes': norma.itens.count(),
    }
    papeis = {coluna['papel'] for coluna in colunas}
    if not {'codigo', 'titulo'} <= papeis:
        pendencia = 'Indique quais colunas são o código do item e a pergunta/título.'
        if dry_run:
            return {**resumo, 'pendencia': pendencia}
        raise NormaItensImportError(pendencia)

    analise = _montar_nos(rows, linha, colunas)
    nos = analise['nos']
    perguntas = [no for no in nos if not no['filhos']]
    por_classificacao = _contar(no['classificacao'] for no in perguntas if no['classificacao'])
    por_area = _contar(area for no in perguntas for area in no['areas'])
    classificacoes_novas = sorted(sigla for sigla in por_classificacao if norma.classificacao_por_sigla(sigla) is None)
    areas_novas = sorted(set(por_area) - norma.siglas_areas())
    resumo.update(
        total=len(nos),
        requisitos=sum(1 for no in nos if no['parent'] is None),
        orientativos=len(nos) - len(perguntas),
        perguntas=len(perguntas),
        porClassificacao=por_classificacao,
        classificacoesNovas=classificacoes_novas,
        porArea=por_area,
        areasNovas=areas_novas,
        linhasIgnoradas=analise['ignoradas'],
        avisos=analise['avisos'],
    )
    if substituir and RespostaAuditoriaSGQ.objects.filter(item__norma=norma).exists():
        resumo['avisos'] = [_SUBSTITUIR_RESPONDIDA, *resumo['avisos']]
        if not dry_run:
            raise NormaItensImportError(_SUBSTITUIR_RESPONDIDA)
    if dry_run:
        return resumo

    classificacoes = {item['sigla']: item for item in novas_classificacoes or [] if item['sigla'] in classificacoes_novas}
    areas = {item['sigla']: item for item in novas_areas or [] if item['sigla'] in areas_novas}
    faltando = [s for s in classificacoes_novas if s not in classificacoes] + [s for s in areas_novas if s not in areas]
    if faltando:
        raise NormaItensImportError(
            f'Defina as classificações/áreas da planilha que não existem na norma: {", ".join(faltando)}.'
        )

    with transaction.atomic():
        norma.classificacoes = [*norma.classificacoes, *(classificacoes[s] for s in classificacoes_novas)]
        norma.areas = [*norma.areas, *({'sigla': s, 'nome': areas[s]['nome']} for s in areas_novas)]
        norma.mapeamento_importacao = [
            {'cabecalho': _normalizar(coluna['cabecalho']), 'papel': coluna['papel'], 'area': coluna['area']}
            for coluna in colunas
        ]
        norma.save(update_fields=['classificacoes', 'areas', 'mapeamento_importacao'])
        if substituir:
            norma.itens.all().delete()
        ultima_raiz = norma.itens.filter(parent__isnull=True).aggregate(maior=Max('ordem'))['maior']
        proxima_ordem: dict[int | None, int] = {None: 0 if ultima_raiz is None else ultima_raiz + 1}
        criados: list[ItemNormaSGQ] = []
        for no in nos:
            pai = criados[no['parent']] if no['parent'] is not None else None
            chave = no['parent']
            ordem = proxima_ordem.get(chave, 0)
            proxima_ordem[chave] = ordem + 1
            pergunta = not no['filhos']
            classificacao = norma.classificacao_por_sigla(no['classificacao']) if pergunta and no['classificacao'] else None
            criados.append(ItemNormaSGQ.objects.create(
                norma=norma,
                parent=pai,
                ordem=ordem,
                codigo=no['codigo'],
                titulo=no['titulo'],
                descricao=no['descricao'],
                tipo='pergunta' if pergunta else 'orientativo',
                tipo_resposta=tipo_resposta if pergunta else None,
                classificacao=classificacao['sigla'] if classificacao else '',
                areas=no['areas'] if pergunta else [],
                peso=1,
                obrigatorio=bool(classificacao['critica']) if classificacao else True,
            ))
    return resumo


# Cabeçalhos que a importação reconhece. Classificação, Área e as marcações são opcionais:
# ISO 9001 usa código, pergunta e orientação; SASSMAQ também preenche classificação e SS/MA/Q.
COLUNAS_MODELO_ITENS = ('Código', 'Pergunta', 'Orientação', 'Classificação', 'Área', 'SS', 'MA', 'Q')

_EXEMPLO_ISO = (
    ('4', 'Contexto da organização', '', '', '', '', '', ''),
    ('4.1', 'Entender a organização e seu contexto', 'Determinar questões externas e internas relevantes.', '', '', '', '', ''),
    ('4.2', 'Necessidades das partes interessadas', 'Determinar as partes interessadas e seus requisitos.', '', '', '', '', ''),
    ('5', 'Liderança', '', '', '', '', '', ''),
    ('5.1', 'Liderança e comprometimento', 'A alta direção demonstra liderança e comprometimento com o SGQ.', '', '', '', '', ''),
)

_EXEMPLO_SASSMAQ = (
    ('1', 'Sistema de gestão', '', '', '', '', '', ''),
    ('1.1', 'Saúde, segurança e meio ambiente', '', '', '', '', '', ''),
    ('1.1.1', 'Há política de SASSMAQ documentada?', 'Verificar o documento da política.', 'M', '', 'X', '', ''),
    ('1.1.2', 'Os aspectos ambientais são identificados?', 'Checar o levantamento de aspectos.', 'I', '', '', 'X', ''),
    ('1.1.3', 'A qualidade do serviço é monitorada?', '', 'D', '', '', '', 'X'),
)

_INSTRUCOES_MODELO = (
    'A aba Lista é a que o sistema importa. As abas de exemplo não entram na importação: copie as linhas para a Lista ou apague o que não for usar.',
    'Código e Pergunta são obrigatórios. A hierarquia vem do código (4 › 4.1 › 4.1.1 ou 1 › 1.1 › 1.1.1a). Item com filhos vira texto orientativo; item sem filhos vira pergunta.',
    'Orientação é opcional.',
    'Classificação é a sigla da pergunta (SASSMAQ: M, I ou D). Na ISO 9001 pode ficar em branco.',
    'Área recebe uma sigla na própria célula, quando a norma não separa as áreas em colunas.',
    'SS, MA e Q são marcações: preencha com X quando a pergunta pertence àquela área. Apague as colunas que a norma não usa ou inclua outra sigla no cabeçalho (ex.: QL).',
    'Na análise, confira o papel de cada coluna. Se a área ainda não existir na norma, escolha Área (marcação) e informe a sigla. O mesmo arquivo serve para ISO 9001, SASSMAQ e outras normas.',
)


def modelo_importacao_itens() -> bytes:
    """Planilha padrão dos requisitos. A primeira aba é a lista vazia; as demais são exemplo e instrução."""
    wb = openpyxl.Workbook()
    lista = wb.active
    lista.title = 'Lista'
    _planilha_modelo(lista, ())

    iso = wb.create_sheet('Exemplo ISO 9001')
    _planilha_modelo(iso, _EXEMPLO_ISO)

    sassmaq = wb.create_sheet('Exemplo SASSMAQ')
    _planilha_modelo(sassmaq, _EXEMPLO_SASSMAQ)

    ajuda = wb.create_sheet('Instruções')
    ajuda['A1'] = 'Como preencher a lista de verificação'
    ajuda['A1'].font = Font(bold=True, size=14, color='0F172A')
    for indice, texto in enumerate(_INSTRUCOES_MODELO, start=3):
        ajuda.cell(indice, 1, texto).alignment = Alignment(wrap_text=True, vertical='top')
        ajuda.row_dimensions[indice].height = 32
    ajuda.column_dimensions['A'].width = 120

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _planilha_modelo(ws, linhas: tuple[tuple, ...]) -> None:
    cabecalho = PatternFill('solid', fgColor='0F76C4')
    fonte = Font(color='FFFFFF', bold=True, size=11)
    alinhamento = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.append(list(COLUNAS_MODELO_ITENS))
    for coluna in range(1, len(COLUNAS_MODELO_ITENS) + 1):
        celula = ws.cell(1, coluna)
        celula.fill = cabecalho
        celula.font = fonte
        celula.alignment = alinhamento
    for linha in linhas:
        ws.append(list(linha))
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:H{max(1, len(linhas) + 1)}'
    ws.row_dimensions[1].height = 22
    for letra, largura in zip('ABCDEFGH', (16, 52, 56, 16, 12, 8, 8, 8), strict=True):
        ws.column_dimensions[letra].width = largura
