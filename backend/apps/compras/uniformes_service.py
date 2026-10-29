"""Regras do Controle de Uniformes.

Toda alteração de saldo passa por `_movimentar`, que grava uma `MovimentacaoUniforme`
com o saldo resultante. Nenhum registro de histórico é alterado ou apagado: correções
são feitas com novos lançamentos (ajuste/descarte)."""
from datetime import date
from io import BytesIO

from django.db import transaction
from django.db.models import F, Q, Sum

from apps.audit.services import record_audit

from .models import (
    CONDICAO_CHOICES,
    CONDICAO_NOVO,
    CONDICAO_USADO,
    MOV_AJUSTE,
    MOV_DESCARTE,
    MOV_DEVOLUCAO,
    MOV_ENTRADA,
    MOV_ENTREGA,
    MOV_PERDA,
    Colaborador,
    DevolucaoUniforme,
    DevolucaoUniformeItem,
    EntregaUniforme,
    EntregaUniformeItem,
    MovimentacaoUniforme,
    Uniforme,
)

CONDICOES_VALIDAS = {c for c, _ in CONDICAO_CHOICES}
FUNCAO_OCULTAR_OBSERVACOES_AJUSTES = 'ocultar-observacoes-ajustes'


def oculta_observacoes_ajustes(user) -> bool:
    """Restrição marcada no usuário: esconde observação e ajustes do histórico."""
    if not getattr(user, 'is_authenticated', False):
        return False
    return FUNCAO_OCULTAR_OBSERVACOES_AJUSTES in ((getattr(user, 'funcoes', None) or {}).get('Compras') or [])


class UniformeError(ValueError):
    pass


def _usuario_nome(user) -> str:
    if not user or not getattr(user, 'is_authenticated', False):
        return ''
    return (getattr(user, 'name', '') or user.get_full_name() or user.username or '').strip()


def _parse_data(value, campo: str) -> date:
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value or '').strip())
    except ValueError:
        raise UniformeError(f'Informe uma {campo} válida.')


def _parse_int(value, campo: str, *, minimo: int = 0) -> int:
    try:
        parsed = int(value or 0)
    except (TypeError, ValueError):
        raise UniformeError(f'{campo} inválida.')
    if parsed < minimo:
        raise UniformeError(f'{campo} deve ser maior ou igual a {minimo}.')
    return parsed


def _parse_condicao(value) -> str:
    condicao = str(value or '').strip().lower()
    if condicao not in CONDICOES_VALIDAS:
        raise UniformeError('Condição inválida: use "novo" ou "usado".')
    return condicao


def _campo_saldo(condicao: str) -> str:
    return 'qtd_nova' if condicao == CONDICAO_NOVO else 'qtd_usada'


def _movimentar(
    uniforme: Uniforme,
    *,
    tipo: str,
    condicao: str,
    quantidade: int,
    variacao: int,
    data: date,
    user,
    observacao: str = '',
    colaborador: Colaborador | None = None,
    colaborador_nome: str = '',
    entrega: EntregaUniforme | None = None,
    devolucao: DevolucaoUniforme | None = None,
) -> MovimentacaoUniforme:
    """Aplica `variacao` ao saldo da condição (uniforme já travado com select_for_update)."""
    if variacao:
        campo = _campo_saldo(condicao)
        novo_saldo = getattr(uniforme, campo) + variacao
        if novo_saldo < 0:
            rotulo = 'novo' if condicao == CONDICAO_NOVO else 'usado'
            raise UniformeError(
                f'Saldo {rotulo} insuficiente para "{uniforme}". Disponível: {getattr(uniforme, campo)}.'
            )
        setattr(uniforme, campo, novo_saldo)
        uniforme.save(update_fields=[campo])

    return MovimentacaoUniforme.objects.create(
        tipo=tipo,
        uniforme=uniforme,
        uniforme_nome=uniforme.nome,
        uniforme_tamanho=uniforme.tamanho,
        condicao=condicao,
        quantidade=quantidade,
        variacao_estoque=variacao,
        saldo_novo=uniforme.qtd_nova,
        saldo_usado=uniforme.qtd_usada,
        data=data,
        colaborador=colaborador,
        colaborador_nome=colaborador_nome or (colaborador.nome if colaborador else ''),
        entrega=entrega,
        devolucao=devolucao,
        observacao=observacao,
        usuario=user if getattr(user, 'is_authenticated', False) else None,
        usuario_nome=_usuario_nome(user),
    )


def _travar(uniforme_id) -> Uniforme:
    uniforme = Uniforme.objects.select_for_update().filter(pk=uniforme_id).first()
    if not uniforme:
        raise UniformeError('Uniforme não encontrado.')
    return uniforme


# ─── Estoque ────────────────────────────────────────────────────────────────

def registrar_entrada(uniforme_id, *, condicao, quantidade, data, observacao='', user) -> MovimentacaoUniforme:
    condicao = _parse_condicao(condicao)
    quantidade = _parse_int(quantidade, 'Quantidade', minimo=1)
    data = _parse_data(data, 'data')
    with transaction.atomic():
        uniforme = _travar(uniforme_id)
        mov = _movimentar(
            uniforme, tipo=MOV_ENTRADA, condicao=condicao, quantidade=quantidade, variacao=quantidade,
            data=data, user=user, observacao=(observacao or '').strip(),
        )
    record_audit(user, 'compras.uniforme.entrada', f'Entrada de {quantidade} {condicao}(s) em "{uniforme}".')
    return mov


def registrar_ajuste(uniforme_id, *, condicao, quantidade, observacao, user) -> MovimentacaoUniforme:
    condicao = _parse_condicao(condicao)
    try:
        variacao = int(quantidade)
    except (TypeError, ValueError):
        raise UniformeError('Quantidade do ajuste inválida.')
    if variacao == 0:
        raise UniformeError('O ajuste deve ter quantidade diferente de zero.')
    observacao = (observacao or '').strip()
    if not observacao:
        raise UniformeError('Informe a justificativa do ajuste.')
    with transaction.atomic():
        uniforme = _travar(uniforme_id)
        mov = _movimentar(
            uniforme, tipo=MOV_AJUSTE, condicao=condicao, quantidade=abs(variacao), variacao=variacao,
            data=date.today(), user=user, observacao=observacao,
        )
    record_audit(user, 'compras.uniforme.ajuste', f'Ajuste de {variacao:+d} {condicao}(s) em "{uniforme}": {observacao}')
    return mov


def registrar_descarte(uniforme_id, *, condicao, quantidade, observacao, user) -> MovimentacaoUniforme:
    condicao = _parse_condicao(condicao)
    quantidade = _parse_int(quantidade, 'Quantidade', minimo=1)
    observacao = (observacao or '').strip()
    if not observacao:
        raise UniformeError('Informe o motivo do descarte.')
    with transaction.atomic():
        uniforme = _travar(uniforme_id)
        mov = _movimentar(
            uniforme, tipo=MOV_DESCARTE, condicao=condicao, quantidade=quantidade, variacao=-quantidade,
            data=date.today(), user=user, observacao=observacao,
        )
    record_audit(user, 'compras.uniforme.descarte', f'Descarte de {quantidade} {condicao}(s) de "{uniforme}": {observacao}')
    return mov


# ─── Entrega ────────────────────────────────────────────────────────────────

def registrar_entrega(*, colaborador_id, data_entrega, observacoes='', itens, user) -> EntregaUniforme:
    colaborador = Colaborador.objects.select_related('setor').filter(pk=colaborador_id).first()
    if not colaborador:
        raise UniformeError('Selecione o colaborador que vai receber os uniformes.')
    data_entrega = _parse_data(data_entrega, 'data de entrega')
    if not itens:
        raise UniformeError('Adicione ao menos um uniforme à entrega.')

    linhas = []
    vistos = set()
    for linha in itens:
        uniforme_id = str(linha.get('uniformeId') or '').strip()
        condicao = _parse_condicao(linha.get('condicao'))
        quantidade = _parse_int(linha.get('quantidade'), 'Quantidade', minimo=1)
        if not uniforme_id:
            raise UniformeError('Selecione o uniforme em todas as linhas.')
        chave = (uniforme_id, condicao)
        if chave in vistos:
            raise UniformeError('O mesmo uniforme e condição aparecem mais de uma vez na entrega.')
        vistos.add(chave)
        linhas.append((uniforme_id, condicao, quantidade))

    with transaction.atomic():
        entrega = EntregaUniforme.objects.create(
            colaborador=colaborador,
            colaborador_nome=colaborador.nome,
            colaborador_cpf=colaborador.cpf,
            colaborador_cargo=colaborador.cargo,
            colaborador_setor=colaborador.setor.nome if colaborador.setor_id else '',
            data_entrega=data_entrega,
            observacoes=(observacoes or '').strip(),
            usuario=user if getattr(user, 'is_authenticated', False) else None,
            usuario_nome=_usuario_nome(user),
        )
        for uniforme_id, condicao, quantidade in linhas:
            uniforme = _travar(uniforme_id)
            EntregaUniformeItem.objects.create(
                entrega=entrega,
                uniforme=uniforme,
                uniforme_nome=uniforme.nome,
                uniforme_codigo=uniforme.codigo,
                uniforme_tamanho=uniforme.tamanho,
                condicao=condicao,
                quantidade=quantidade,
            )
            _movimentar(
                uniforme, tipo=MOV_ENTREGA, condicao=condicao, quantidade=quantidade, variacao=-quantidade,
                data=data_entrega, user=user, colaborador=colaborador, entrega=entrega,
                observacao=f'Termo {entrega.numero}',
            )

    record_audit(
        user, 'compras.uniforme.entrega',
        f'Entrega {entrega.numero} para "{colaborador.nome}" com {len(linhas)} item(ns).',
    )
    return entrega


# ─── Devolução ──────────────────────────────────────────────────────────────

def registrar_devolucao(entrega_id, *, data_devolucao, observacoes='', itens, user) -> DevolucaoUniforme:
    data_devolucao = _parse_data(data_devolucao, 'data de devolução')
    observacoes = (observacoes or '').strip()
    if not itens:
        raise UniformeError('Informe ao menos um item na devolução.')

    with transaction.atomic():
        entrega = EntregaUniforme.objects.select_for_update().filter(pk=entrega_id).first()
        if not entrega:
            raise UniformeError('Entrega não encontrada.')
        itens_entrega = {str(i.pk): i for i in entrega.itens.select_for_update()}

        linhas = []
        for linha in itens:
            item = itens_entrega.get(str(linha.get('entregaItemId') or ''))
            if not item:
                raise UniformeError('Item não pertence a esta entrega.')
            boa = _parse_int(linha.get('qtdBoa'), 'Quantidade em boas condições')
            danificada = _parse_int(linha.get('qtdDanificada'), 'Quantidade danificada')
            sem_condicoes = _parse_int(linha.get('qtdSemCondicoes'), 'Quantidade sem condições')
            perdida = _parse_int(linha.get('qtdPerdida'), 'Quantidade perdida')
            total = boa + danificada + sem_condicoes + perdida
            if total == 0:
                continue
            if total > item.pendente:
                raise UniformeError(
                    f'"{item.uniforme_nome} ({item.uniforme_tamanho})": informado {total}, '
                    f'mas só {item.pendente} peça(s) estão pendentes de devolução.'
                )
            linhas.append((item, boa, danificada, sem_condicoes, perdida))

        if not linhas:
            raise UniformeError('Informe a quantidade devolvida ou perdida de ao menos um item.')
        if any(perdida for *_, perdida in linhas) and not observacoes:
            raise UniformeError('Informe na observação a justificativa das peças perdidas ou não devolvidas.')

        devolucao = DevolucaoUniforme.objects.create(
            entrega=entrega,
            data_devolucao=data_devolucao,
            observacoes=observacoes,
            usuario=user if getattr(user, 'is_authenticated', False) else None,
            usuario_nome=_usuario_nome(user),
        )
        ref = f'Devolução do termo {entrega.numero}'
        for item, boa, danificada, sem_condicoes, perdida in linhas:
            DevolucaoUniformeItem.objects.create(
                devolucao=devolucao, entrega_item=item, qtd_boa=boa, qtd_danificada=danificada,
                qtd_sem_condicoes=sem_condicoes, qtd_perdida=perdida,
            )
            EntregaUniformeItem.objects.filter(pk=item.pk).update(
                qtd_devolvida=F('qtd_devolvida') + boa,
                qtd_descartada=F('qtd_descartada') + danificada + sem_condicoes,
                qtd_perdida=F('qtd_perdida') + perdida,
            )

            uniforme = _travar(item.uniforme_id) if item.uniforme_id else None
            if uniforme is None:
                if boa:
                    raise UniformeError(
                        f'O cadastro de "{item.uniforme_nome} ({item.uniforme_tamanho})" foi removido; '
                        'não é possível devolver ao estoque.'
                    )
                continue
            comum = dict(
                data=data_devolucao, user=user, colaborador=entrega.colaborador,
                colaborador_nome=entrega.colaborador_nome, entrega=entrega, devolucao=devolucao,
            )
            if boa:
                _movimentar(
                    uniforme, tipo=MOV_DEVOLUCAO, condicao=CONDICAO_USADO, quantidade=boa, variacao=boa,
                    observacao=f'{ref} — usada em boas condições', **comum,
                )
            descarte = danificada + sem_condicoes
            if descarte:
                detalhe = ', '.join(filter(None, [
                    f'{danificada} danificada(s)' if danificada else '',
                    f'{sem_condicoes} sem condições de uso' if sem_condicoes else '',
                ]))
                _movimentar(
                    uniforme, tipo=MOV_DESCARTE, condicao=CONDICAO_USADO, quantidade=descarte, variacao=0,
                    observacao=f'{ref} — {detalhe}', **comum,
                )
            if perdida:
                _movimentar(
                    uniforme, tipo=MOV_PERDA, condicao=item.condicao, quantidade=perdida, variacao=0,
                    observacao=f'{ref} — {observacoes}', **comum,
                )

    record_audit(
        user, 'compras.uniforme.devolucao',
        f'Devolução registrada no termo {entrega.numero} ({entrega.colaborador_nome}).',
    )
    return devolucao


def registrar_devolucao_total(colaborador_id, *, data_devolucao, observacoes='', itens, user) -> list[DevolucaoUniforme]:
    """Devolução total: dá destino a todas as peças pendentes do colaborador, em todos os termos."""
    observacoes = (observacoes or '').strip()
    colaborador = Colaborador.objects.filter(pk=colaborador_id).first()
    if not colaborador:
        raise UniformeError('Colaborador não encontrado.')

    with transaction.atomic():
        pendentes = list(
            EntregaUniformeItem.objects.select_for_update()
            .filter(entrega__colaborador=colaborador)
            .annotate(pend=F('quantidade') - F('qtd_devolvida') - F('qtd_descartada') - F('qtd_perdida'))
            .filter(pend__gt=0)
        )
        if not pendentes:
            raise UniformeError('Este colaborador não tem uniformes pendentes de devolução.')

        informados = {str(linha.get('entregaItemId') or ''): linha for linha in (itens or [])}
        por_entrega: dict[int, list[dict]] = {}
        for item in pendentes:
            linha = informados.get(str(item.pk))
            total = 0
            if linha:
                total = sum(
                    _parse_int(linha.get(campo), 'Quantidade')
                    for campo in ('qtdBoa', 'qtdDanificada', 'qtdSemCondicoes', 'qtdPerdida')
                )
            if total != item.pend:
                raise UniformeError(
                    f'"{item.uniforme_nome} ({item.uniforme_tamanho})": informe o destino de todas as '
                    f'{item.pend} peça(s) — na devolução total nada pode ficar pendente.'
                )
            por_entrega.setdefault(item.entrega_id, []).append(linha)

        if not observacoes and any(_parse_int(l.get('qtdPerdida'), 'Quantidade') for ls in por_entrega.values() for l in ls):
            raise UniformeError('Informe na observação a justificativa das peças perdidas ou não devolvidas.')
        texto = f'Devolução total — {observacoes}' if observacoes else 'Devolução total'
        devolucoes = [
            registrar_devolucao(entrega_id, data_devolucao=data_devolucao, observacoes=texto, itens=linhas, user=user)
            for entrega_id, linhas in sorted(por_entrega.items())
        ]

    record_audit(
        user, 'compras.uniforme.devolucao_total',
        f'Devolução total de {colaborador.nome}: {len(devolucoes)} termo(s).',
    )
    return devolucoes


# ─── Consultas ──────────────────────────────────────────────────────────────

def totais_por_uniforme() -> dict[int, dict[str, int]]:
    """Peças destinadas (em posse), perdidas e descartadas por uniforme."""
    totais: dict[int, dict[str, int]] = {}

    def _slot(uid):
        return totais.setdefault(uid, {'destinados': 0, 'perdidos': 0, 'descartados': 0})

    pendentes = (
        EntregaUniformeItem.objects.filter(uniforme__isnull=False)
        .values('uniforme_id')
        .annotate(total=Sum(F('quantidade') - F('qtd_devolvida') - F('qtd_descartada') - F('qtd_perdida')))
    )
    for row in pendentes:
        _slot(row['uniforme_id'])['destinados'] = row['total'] or 0

    for tipo, chave in ((MOV_PERDA, 'perdidos'), (MOV_DESCARTE, 'descartados')):
        rows = (
            MovimentacaoUniforme.objects.filter(tipo=tipo, uniforme__isnull=False)
            .values('uniforme_id').annotate(total=Sum('quantidade'))
        )
        for row in rows:
            _slot(row['uniforme_id'])[chave] = row['total'] or 0
    return totais


def resumo() -> dict[str, int]:
    estoque = Uniforme.objects.aggregate(novo=Sum('qtd_nova'), usado=Sum('qtd_usada'))
    totais = totais_por_uniforme().values()
    baixo = Uniforme.objects.filter(
        Q(estoque_minimo__gt=F('qtd_nova') + F('qtd_usada')) | Q(qtd_nova=0, qtd_usada=0)
    ).count()
    return {
        'estoqueNovo': estoque['novo'] or 0,
        'estoqueUsado': estoque['usado'] or 0,
        'destinados': sum(t['destinados'] for t in totais),
        'perdidos': sum(t['perdidos'] for t in totais),
        'descartados': sum(t['descartados'] for t in totais),
        'estoqueBaixo': baixo,
        'modelos': Uniforme.objects.count(),
    }


def planilha_estoque() -> bytes:
    """Estoque por tamanho e as peças que ainda estão com colaboradores."""
    import openpyxl

    totais = totais_por_uniforme()
    wb = openpyxl.Workbook()
    estoque = wb.active
    estoque.title = 'Estoque'
    _cabecalho(estoque, [
        'Item', 'Código', 'Tamanho', 'Novas', 'Usadas', 'Disponível',
        'Estoque mínimo', 'Com colaboradores', 'Situação',
    ])
    for uniforme in Uniforme.objects.order_by('nome', 'tamanho'):
        novas = uniforme.qtd_nova
        usadas = uniforme.qtd_usada
        disponivel = novas + usadas
        com_colaboradores = totais.get(uniforme.pk, {}).get('destinados', 0)
        estoque.append([
            uniforme.nome,
            uniforme.codigo,
            uniforme.tamanho,
            novas,
            usadas,
            disponivel,
            uniforme.estoque_minimo,
            com_colaboradores,
            _situacao_estoque(disponivel, uniforme.estoque_minimo),
        ])
    _ajustar_planilha(estoque)

    posse = wb.create_sheet('Com colaboradores')
    _cabecalho(posse, [
        'Colaborador', 'CPF', 'Cargo', 'Setor', 'Termo', 'Data da entrega',
        'Item', 'Código', 'Tamanho', 'Condição', 'Quantidade',
    ])
    condicao = dict(CONDICAO_CHOICES)
    itens = (
        EntregaUniformeItem.objects
        .annotate(em_posse=F('quantidade') - F('qtd_devolvida') - F('qtd_descartada') - F('qtd_perdida'))
        .filter(em_posse__gt=0)
        .select_related('entrega')
        .order_by('entrega__colaborador_nome', 'uniforme_nome', 'uniforme_tamanho', 'entrega__data_entrega')
    )
    for item in itens:
        entrega = item.entrega
        posse.append([
            entrega.colaborador_nome,
            _formatar_cpf(entrega.colaborador_cpf),
            entrega.colaborador_cargo,
            entrega.colaborador_setor,
            entrega.numero,
            entrega.data_entrega,
            item.uniforme_nome,
            item.uniforme_codigo,
            item.uniforme_tamanho,
            condicao.get(item.condicao, item.condicao),
            item.em_posse,
        ])
    for linha in posse.iter_rows(min_row=2, min_col=6, max_col=6):
        for celula in linha:
            celula.number_format = 'DD/MM/YYYY'
    _ajustar_planilha(posse)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _situacao_estoque(disponivel: int, minimo: int) -> str:
    if disponivel == 0:
        return 'Esgotado'
    if minimo > 0 and disponivel < minimo:
        return 'Estoque baixo'
    return 'Disponível'


def _formatar_cpf(valor: str) -> str:
    digitos = ''.join(caractere for caractere in (valor or '') if caractere.isdigit())
    if len(digitos) != 11:
        return valor or ''
    return f'{digitos[:3]}.{digitos[3:6]}.{digitos[6:9]}-{digitos[9:]}'


def _cabecalho(planilha, colunas: list[str]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    planilha.append(colunas)
    preenchimento = PatternFill('solid', fgColor='118CC4')
    fonte = Font(bold=True, color='FFFFFF')
    for celula in planilha[1]:
        celula.fill = preenchimento
        celula.font = fonte
        celula.alignment = Alignment(horizontal='center')


def _ajustar_planilha(planilha) -> None:
    from openpyxl.utils import get_column_letter

    planilha.freeze_panes = 'A2'
    planilha.auto_filter.ref = planilha.dimensions
    for indice, coluna in enumerate(planilha.columns, start=1):
        largura = max(len(str(celula.value or '')) for celula in coluna)
        planilha.column_dimensions[get_column_letter(indice)].width = min(max(largura + 2, 12), 36)
