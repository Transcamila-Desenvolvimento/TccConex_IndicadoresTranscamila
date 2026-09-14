import base64
from django.contrib.auth import get_user_model
from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from rest_framework.test import APIClient
from unittest.mock import patch

from .models import ClienteComercial, PropostaComercialDraft

User = get_user_model()

HEADERS = {'HTTP_X_PROTHON_ENVIRONMENT': 'Comercial'}
PDF_BASE64 = base64.b64encode(b'%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF\n').decode()

PAYLOAD = {
    'tipoPessoa': 'J',
    'cnpj': '00.000.000/0001-91',
    'razaoSocial': 'Empresa Teste Ltda',
    'nomeFantasia': 'Empresa Teste',
    'municipio': 'curitiba',
    'uf': 'pr',
    'situacao': 'cliente',
}


class ClienteComercialTests(TestCase):
    def setUp(self):
        self.api = APIClient()
        self.admin = User.objects.create_user(
            username='comercial.admin',
            password='test123',
            name='Admin Comercial',
            role_id='1',
            status='ativo',
            environments=['Comercial'],
        )
        self.operador = User.objects.create_user(
            username='comercial.op',
            password='test123',
            name='Operador Comercial',
            role_id='2',
            status='ativo',
            environments=['Comercial'],
            funcoes={'Comercial': ['gerenciar-clientes']},
        )
        self.leitura = User.objects.create_user(
            username='comercial.view',
            password='test123',
            name='Leitura Comercial',
            role_id='2',
            status='ativo',
            environments=['Comercial'],
        )

    def _auth(self, user):
        self.api.force_authenticate(user=user)

    def _publicar_tabela_distribuicao(self, cliente_id, nome='Tabela Dist'):
        created = self.api.post(
            '/api/comercial/tabela-frete/',
            {
                'nome': nome,
                'tipo': 'distribuicao',
                'clienteIds': [cliente_id],
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(created.status_code, 201, created.content)
        published = self.api.post(
            f'/api/comercial/tabela-frete/{created.json()["id"]}/publicar/',
            **HEADERS,
        )
        self.assertEqual(published.status_code, 200, published.content)
        return created.json()['id']

    def test_list_requires_auth(self):
        response = self.api.get('/api/comercial/clientes/')
        self.assertIn(response.status_code, (401, 403))

    def test_summary_ok(self):
        self._auth(self.admin)
        response = self.api.get('/api/comercial/summary/', **HEADERS)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['environment'], 'Comercial')

    def test_admin_cria_e_lista_cliente(self):
        self._auth(self.admin)
        created = self.api.post('/api/comercial/clientes/', PAYLOAD, format='json', **HEADERS)
        self.assertEqual(created.status_code, 201, created.content)
        body = created.json()
        self.assertEqual(body['cnpj'], '00.000.000/0001-91')
        self.assertEqual(body['razaoSocial'], 'EMPRESA TESTE LTDA')
        self.assertEqual(body['municipio'], 'Curitiba')
        self.assertEqual(body['uf'], 'PR')
        self.assertEqual(body['situacao'], 'cliente')
        self.assertEqual(body['tipoPessoa'], 'J')
        self.assertEqual(body.get('responsavel', ''), '')
        self.assertEqual(body['compatibilidade'], 'nao_analisado')

        listed = self.api.get('/api/comercial/clientes/', **HEADERS)
        self.assertEqual(listed.status_code, 200)
        self.assertEqual(listed.json()['count'], 1)
        self.assertEqual(len(listed.json()['results']), 1)

    def test_cria_cliente_sempre_nao_analisado(self):
        self._auth(self.admin)
        created = self.api.post(
            '/api/comercial/clientes/',
            {**PAYLOAD, 'compatibilidade': 'compativel'},
            format='json',
            **HEADERS,
        )
        self.assertEqual(created.status_code, 201, created.content)
        self.assertEqual(created.json()['compatibilidade'], 'nao_analisado')

        updated = self.api.patch(
            f'/api/comercial/clientes/{created.json()["id"]}/',
            {'compatibilidade': 'incompativel'},
            format='json',
            **HEADERS,
        )
        self.assertEqual(updated.status_code, 200, updated.content)
        self.assertEqual(updated.json()['compatibilidade'], 'nao_analisado')

    def test_cliente_nao_grava_produtos_no_cadastro(self):
        self._auth(self.admin)
        created = self.api.post(
            '/api/comercial/clientes/',
            {
                **PAYLOAD,
                'produtos': [
                    {
                        'nome': 'Glifosato 480',
                        'fispq': 'https://exemplo.com/fispq.pdf',
                        'numeroOnu': '3082',
                        'classeRisco': '9',
                        'grupoEmbalagem': 'III',
                    },
                ],
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(created.status_code, 201, created.content)
        self.assertEqual(created.json()['produtos'], [])
        self.assertEqual(created.json()['produtosCount'], 0)

    def test_sugestoes_de_produtos_incluem_padrao(self):
        self._auth(self.admin)
        response = self.api.get('/api/comercial/clientes/produtos-sugestoes/', **HEADERS)
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertTrue(any(item['value'] == '9' for item in body['classesRisco']))
        self.assertTrue(any(item['numeroOnu'] == '3082' for item in body['onuComuns']))
        self.assertTrue(any(item['value'] == 'pendente_validacao' for item in body['homologacao']))

    def test_cnpj_duplicado(self):
        self._auth(self.admin)
        first = self.api.post('/api/comercial/clientes/', PAYLOAD, format='json', **HEADERS)
        self.assertEqual(first.status_code, 201, first.content)
        second = self.api.post('/api/comercial/clientes/', PAYLOAD, format='json', **HEADERS)
        self.assertEqual(second.status_code, 400)

    def test_responsavel_do_cliente_preenche_proposta(self):
        self._auth(self.admin)
        cliente = self.api.post(
            '/api/comercial/clientes/',
            {**PAYLOAD, 'responsavel': 'Bianca Silva'},
            format='json',
            **HEADERS,
        )
        self.assertEqual(cliente.status_code, 201, cliente.content)
        self.assertEqual(cliente.json()['responsavel'], 'Bianca Silva')

        proposta = self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'armazenagem',
                'titulo': 'Proposta',
                'clienteId': cliente.json()['id'],
                'att': 'ignorado',
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(proposta.status_code, 201, proposta.content)
        self.assertEqual(proposta.json()['att'], 'Bianca Silva')

    def test_operador_sem_funcao_nao_cria(self):
        self._auth(self.leitura)
        response = self.api.post('/api/comercial/clientes/', PAYLOAD, format='json', **HEADERS)
        self.assertEqual(response.status_code, 403)
        self.assertEqual(ClienteComercial.objects.count(), 0)

    def test_operador_com_funcao_cria(self):
        self._auth(self.operador)
        response = self.api.post('/api/comercial/clientes/', PAYLOAD, format='json', **HEADERS)
        self.assertEqual(response.status_code, 201, response.content)

    def test_busca_paginada(self):
        self._auth(self.admin)
        self.api.post('/api/comercial/clientes/', PAYLOAD, format='json', **HEADERS)
        other = {**PAYLOAD, 'cnpj': '00.000.000/0002-72', 'razaoSocial': 'Outra Empresa'}
        self.api.post('/api/comercial/clientes/', other, format='json', **HEADERS)
        listed = self.api.get('/api/comercial/clientes/?search=outra', **HEADERS)
        self.assertEqual(listed.status_code, 200)
        self.assertEqual(listed.json()['count'], 1)
        self.assertIn('OUTRA', listed.json()['results'][0]['razaoSocial'])

    def test_rejeita_pessoa_fisica(self):
        self._auth(self.admin)
        response = self.api.post(
            '/api/comercial/clientes/',
            {**PAYLOAD, 'tipoPessoa': 'F', 'cnpj': '390.533.447-05'},
            format='json',
            **HEADERS,
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(ClienteComercial.objects.count(), 0)

    def test_filtra_situacao(self):
        self._auth(self.admin)
        self.api.post('/api/comercial/clientes/', PAYLOAD, format='json', **HEADERS)
        self.api.post(
            '/api/comercial/clientes/',
            {**PAYLOAD, 'cnpj': '00.000.000/0002-72', 'razaoSocial': 'Prospect Ltda', 'situacao': 'potencial'},
            format='json',
            **HEADERS,
        )
        listed = self.api.get('/api/comercial/clientes/?situacao=potencial', **HEADERS)
        self.assertEqual(listed.status_code, 200)
        self.assertEqual(listed.json()['count'], 1)
        self.assertEqual(listed.json()['results'][0]['situacao'], 'potencial')

    def test_consultar_cnpj_admin(self):
        self._auth(self.admin)
        with patch('apps.comercial.views.consultar_cnpj') as mock_cnpj:
            mock_cnpj.return_value = {
                'cnpj': '00.000.000/0001-91',
                'razaoSocial': 'EMPRESA TESTE LTDA',
                'nomeFantasia': 'EMPRESA TESTE',
                'municipio': 'Curitiba',
                'uf': 'PR',
            }
            response = self.api.get(
                '/api/comercial/clientes/consultar-cnpj/?cnpj=00000000000191',
                **HEADERS,
            )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['razaoSocial'], 'EMPRESA TESTE LTDA')
        self.assertEqual(response.data['uf'], 'PR')

    def test_cria_e_filtra_proposta(self):
        self._auth(self.admin)
        cliente = self.api.post('/api/comercial/clientes/', PAYLOAD, format='json', **HEADERS)
        self.assertEqual(cliente.status_code, 201, cliente.content)
        created = self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'transporte_rodoviario',
                'titulo': 'Tabela SP-PR',
                'clienteId': cliente.json()['id'],
                'valorEstimado': '1500.50',
                'status': 'rascunho',
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(created.status_code, 201, created.content)
        self.assertEqual(created.json()['tipo'], 'transporte_rodoviario')
        self.assertEqual(created.json()['titulo'], 'Tabela SP-PR')
        self.assertEqual(created.json()['clienteNome'], 'EMPRESA TESTE LTDA')

        self.api.post(
            '/api/comercial/propostas/',
            {'tipo': 'armazenagem', 'titulo': 'Palete mensal', 'status': 'enviada'},
            format='json',
            **HEADERS,
        )
        listed = self.api.get('/api/comercial/propostas/?tipo=transporte_rodoviario', **HEADERS)
        self.assertEqual(listed.status_code, 200)
        self.assertEqual(listed.json()['count'], 1)
        self.assertEqual(listed.json()['results'][0]['tipo'], 'transporte_rodoviario')

    def test_dashboard_propostas(self):
        self._auth(self.admin)
        cliente = self.api.post('/api/comercial/clientes/', PAYLOAD, format='json', **HEADERS)
        self.assertEqual(cliente.status_code, 201, cliente.content)
        self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'transporte_rodoviario',
                'titulo': 'Frete',
                'clienteId': cliente.json()['id'],
                'status': 'rascunho',
            },
            format='json',
            **HEADERS,
        )
        self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'armazenagem',
                'titulo': 'CD',
                'clienteId': cliente.json()['id'],
                'status': 'enviada',
            },
            format='json',
            **HEADERS,
        )
        response = self.api.get('/api/comercial/propostas/dashboard/', **HEADERS)
        self.assertEqual(response.status_code, 200, response.content)
        body = response.json()
        self.assertEqual(body['total'], 2)
        self.assertEqual(body['porStatus']['rascunho'], 1)
        self.assertEqual(body['porStatus']['enviada'], 1)
        self.assertEqual(body['porTipo']['transporte_rodoviario'], 1)
        self.assertEqual(body['porTipo']['armazenagem'], 1)
        self.assertEqual(len(body['recentes']), 2)
        self.assertEqual(body['recentes'][0]['clienteNome'], 'EMPRESA TESTE LTDA')
        self.assertEqual(len(body['porMes']), 6)
        self.assertIn('criadas', body['porMes'][-1])
        filtrado = self.api.get(
            f'/api/comercial/propostas/dashboard/?cliente={cliente.json()["id"]}',
            **HEADERS,
        )
        self.assertEqual(filtrado.status_code, 200, filtrado.content)
        self.assertEqual(filtrado.json()['total'], 2)
        vazio = self.api.get(
            '/api/comercial/propostas/dashboard/?cliente=999999',
            **HEADERS,
        )
        self.assertEqual(vazio.status_code, 200, vazio.content)
        self.assertEqual(vazio.json()['total'], 0)

    def test_proposta_modalidades_transporte(self):
        self._auth(self.admin)
        sem_cliente = self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'transporte_rodoviario',
                'titulo': 'Com modalidades',
                'incluiTransferencia': True,
                'incluiDistribuicao': True,
                'status': 'rascunho',
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(sem_cliente.status_code, 400, sem_cliente.content)

        cliente = self.api.post('/api/comercial/clientes/', PAYLOAD, format='json', **HEADERS)
        self.assertEqual(cliente.status_code, 201, cliente.content)
        cliente_id = cliente.json()['id']
        sem_tabela = self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'transporte_rodoviario',
                'clienteId': cliente_id,
                'incluiTransferencia': True,
                'incluiDistribuicao': True,
                'status': 'rascunho',
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(sem_tabela.status_code, 400, sem_tabela.content)
        self.assertIn('tabela', str(sem_tabela.json()).lower())

        self._publicar_tabela_distribuicao(cliente_id)
        created = self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'transporte_rodoviario',
                'clienteId': cliente_id,
                'titulo': 'Com modalidades',
                'incluiTransferencia': True,
                'incluiDistribuicao': True,
                'status': 'rascunho',
                'linhas': [
                    {
                        'origem': 'Maringá - PR',
                        'entrega': 'São Paulo - SP',
                        'veiculo': 'Carreta',
                        'tarifaFrete': '1000.00',
                        'prazoDias': '3 dias úteis',
                    }
                ],
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(created.status_code, 201, created.content)
        self.assertTrue(created.json()['incluiTransferencia'])
        self.assertTrue(created.json()['incluiDistribuicao'])

        armazenagem = self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'armazenagem',
                'titulo': 'Sem modalidades',
                'incluiTransferencia': True,
                'incluiDistribuicao': True,
                'status': 'rascunho',
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(armazenagem.status_code, 201, armazenagem.content)
        self.assertFalse(armazenagem.json()['incluiTransferencia'])
        self.assertFalse(armazenagem.json()['incluiDistribuicao'])

    def test_recusa_transferencia_sem_destinos_preenchidos(self):
        self._auth(self.admin)
        cliente = self.api.post('/api/comercial/clientes/', PAYLOAD, format='json', **HEADERS)
        self.assertEqual(cliente.status_code, 201, cliente.content)
        vazia = self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'transporte_rodoviario',
                'clienteId': cliente.json()['id'],
                'incluiTransferencia': True,
                'status': 'rascunho',
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(vazia.status_code, 400, vazia.content)
        self.assertIn('destinos', str(vazia.json()).lower())

        incompleta = self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'transporte_rodoviario',
                'clienteId': cliente.json()['id'],
                'incluiTransferencia': True,
                'status': 'rascunho',
                'linhas': [{'origem': 'Maringá - PR', 'entrega': '', 'veiculo': 'Carreta'}],
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(incompleta.status_code, 400, incompleta.content)

        preenchida = self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'transporte_rodoviario',
                'clienteId': cliente.json()['id'],
                'incluiTransferencia': True,
                'status': 'rascunho',
                'linhas': [
                    {
                        'origem': 'Maringá - PR',
                        'entrega': 'Curitiba - PR',
                        'veiculo': 'Carreta',
                        'tarifaFrete': '1500.00',
                        'prazoDias': '2 dias úteis',
                    }
                ],
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(preenchida.status_code, 201, preenchida.content)
        self.assertEqual(len(preenchida.json()['linhas']), 1)

    def test_proposta_armazenagem_grava_tabela(self):
        self._auth(self.admin)
        cliente = self.api.post('/api/comercial/clientes/', PAYLOAD, format='json', **HEADERS)
        self.assertEqual(cliente.status_code, 201, cliente.content)
        created = self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'armazenagem',
                'clienteId': cliente.json()['id'],
                'status': 'rascunho',
                'tabelaArmazenagem': {
                    'unidade': 'MT',
                    'itens': [
                        {'rotulo': 'FATURAMENTO MÍNIMO (1)', 'valor': 'R$ 10.000,00'},
                    ],
                    'horaExtraTitulo': 'Hora-extra (7)',
                    'horaExtra': [
                        {'periodo': 'De segunda a sábado', 'valor': 'R$ 20,81/ton'},
                    ],
                    'expediente': 'Expediente do CD: de seg a sex das 08:00 às 17:00h',
                },
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(created.status_code, 201, created.content)
        tabela = created.json()['tabelaArmazenagem']
        self.assertEqual(tabela['unidade'], 'MT')
        self.assertEqual(tabela['itens'][0]['valor'], 'R$ 10.000,00')
        self.assertEqual(tabela['itens'][0]['formato'], 'moeda')
        self.assertEqual(tabela['horaExtra'][0]['periodo'], 'De segunda a sábado')
        self.assertEqual(tabela['horaExtra'][0]['formato'], 'tonelada')

        padrao = self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'armazenagem',
                'clienteId': cliente.json()['id'],
                'titulo': 'Padrão',
                'status': 'rascunho',
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(padrao.status_code, 201, padrao.content)
        tabela_padrao = padrao.json()['tabelaArmazenagem']
        self.assertEqual(tabela_padrao['itens'][0]['rotulo'], 'FATURAMENTO MÍNIMO (1)')
        self.assertEqual(tabela_padrao['itens'][0]['valor'], '-')
        self.assertEqual(tabela_padrao['itens'][3]['formato'], 'percentual')

    def test_proposta_armazenagem_rejeita_tarifa_vazia(self):
        self._auth(self.admin)
        cliente = self.api.post('/api/comercial/clientes/', PAYLOAD, format='json', **HEADERS)
        response = self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'armazenagem',
                'clienteId': cliente.json()['id'],
                'status': 'rascunho',
                'tabelaArmazenagem': {
                    'itens': [
                        {'rotulo': 'FATURAMENTO MÍNIMO (1)', 'valor': ''},
                    ],
                    'horaExtra': [
                        {'periodo': 'De segunda a sábado', 'valor': '-'},
                    ],
                },
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(response.status_code, 400, response.content)
        self.assertIn('tabelaArmazenagem', response.json())

    def test_salvar_proposta_atualiza_catalogo_generalidades(self):
        self._auth(self.admin)
        cliente = self.api.post('/api/comercial/clientes/', PAYLOAD, format='json', **HEADERS)
        self.assertEqual(cliente.status_code, 201, cliente.content)
        cliente_id = cliente.json()['id']
        created = self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'transporte_rodoviario',
                'clienteId': cliente_id,
                'incluiTransferencia': True,
                'status': 'rascunho',
                'linhas': [
                    {
                        'origem': 'Maringá - PR',
                        'entrega': 'Curitiba - PR',
                        'veiculo': 'Carreta',
                        'tarifaFrete': '1500.00',
                        'prazoDias': '2 dias úteis',
                    }
                ],
                'condicoes': [
                    {'rotulo': 'Pedágios', 'valor': 'Incluso no frete', 'tipo': 'frete'},
                    {'rotulo': 'Cubagem', 'valor': 'Alterada na proposta', 'tipo': 'frete'},
                ],
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(created.status_code, 201, created.content)
        catalogo = self.api.get(
            f'/api/comercial/generalidades/?cliente={cliente_id}&tipo=frete',
            **HEADERS,
        )
        self.assertEqual(catalogo.status_code, 200)
        self.assertEqual(catalogo.json()['origem'], 'cliente')
        rotulos = [item['rotulo'] for item in catalogo.json()['items']]
        self.assertEqual(rotulos, ['Pedágios', 'Cubagem'])
        self.assertEqual(catalogo.json()['items'][0]['valor'], 'Incluso no frete')

        atualizada = self.api.patch(
            f'/api/comercial/propostas/{created.json()["id"]}/',
            {
                'condicoes': [
                    {'rotulo': 'Pedágios', 'valor': 'Por trecho', 'tipo': 'frete'},
                ],
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(atualizada.status_code, 200, atualizada.content)
        catalogo = self.api.get(
            f'/api/comercial/generalidades/?cliente={cliente_id}&tipo=frete',
            **HEADERS,
        )
        self.assertEqual(catalogo.json()['items'][0]['valor'], 'Por trecho')
        self.assertEqual(len(catalogo.json()['items']), 1)

    def test_lista_propostas_ordena_por_data_e_vencimento(self):
        self._auth(self.admin)
        antiga = self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'armazenagem',
                'titulo': 'Antiga',
                'dataProposta': '2026-01-01',
                'validade': '30 dias',
                'status': 'rascunho',
            },
            format='json',
            **HEADERS,
        )
        recente = self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'armazenagem',
                'titulo': 'Recente',
                'dataProposta': '2026-06-01',
                'validade': '10 dias',
                'status': 'enviada',
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(antiga.status_code, 201, antiga.content)
        self.assertEqual(recente.status_code, 201, recente.content)

        por_criacao = self.api.get('/api/comercial/propostas/?ordering=data_criacao_desc', **HEADERS)
        self.assertEqual(por_criacao.status_code, 200)
        titulos_criacao = [item['titulo'] for item in por_criacao.json()['results']]
        self.assertEqual(titulos_criacao[:2], ['Recente', 'Antiga'])

        por_vencimento = self.api.get('/api/comercial/propostas/?ordering=vencimento_asc', **HEADERS)
        self.assertEqual(por_vencimento.status_code, 200)
        titulos_vencimento = [item['titulo'] for item in por_vencimento.json()['results']]
        self.assertEqual(titulos_vencimento[:2], ['Antiga', 'Recente'])

    def test_numera_propostas_sequenciais_por_ano(self):
        self._auth(self.admin)
        primeira = self.api.post(
            '/api/comercial/propostas/',
            {'tipo': 'armazenagem', 'titulo': 'Primeira', 'dataProposta': '2026-01-10', 'status': 'rascunho'},
            format='json',
            **HEADERS,
        )
        segunda = self.api.post(
            '/api/comercial/propostas/',
            {'tipo': 'armazenagem', 'titulo': 'Segunda', 'dataProposta': '2026-06-01', 'status': 'enviada'},
            format='json',
            **HEADERS,
        )
        outra_ano = self.api.post(
            '/api/comercial/propostas/',
            {'tipo': 'armazenagem', 'titulo': 'Outro ano', 'dataProposta': '2025-12-01', 'status': 'rascunho'},
            format='json',
            **HEADERS,
        )
        self.assertEqual(primeira.status_code, 201, primeira.content)
        self.assertEqual(primeira.json()['numeroIdentificacao'], '001-2026')
        self.assertEqual(segunda.json()['numeroIdentificacao'], '002-2026')
        self.assertEqual(outra_ano.json()['numeroIdentificacao'], '001-2025')

        atualizada = self.api.patch(
            f'/api/comercial/propostas/{primeira.json()["id"]}/',
            {'titulo': 'Primeira revisada'},
            format='json',
            **HEADERS,
        )
        self.assertEqual(atualizada.status_code, 200, atualizada.content)
        self.assertEqual(atualizada.json()['numeroIdentificacao'], '001-2026')

        localizada = self.api.get('/api/comercial/propostas/?search=002-2026', **HEADERS)
        self.assertEqual(localizada.status_code, 200)
        self.assertEqual(localizada.json()['count'], 1)
        self.assertEqual(localizada.json()['results'][0]['titulo'], 'Segunda')

    def test_cria_proposta_frete_com_tabela_e_total(self):
        self._auth(self.admin)
        created = self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'transporte_rodoviario',
                'titulo': 'PROPOSTA COMERCIAL DE FRETE',
                'subtitulo': 'Porto de Paranaguá – Logística de Contêineres',
                'revisao': '01',
                'dataProposta': '2026-03-25',
                'clienteNome': 'M - OCEAN SHIPPING',
                'propostaReferente': 'Porto Paranaguá',
                'responsavel': 'Marco Aurelio Coradassi',
                'att': 'Bianca',
                'validade': '5 dias',
                'linhas': [
                    {
                        'origem': 'Paranaguá - PR',
                        'entrega': 'Guarapuava - PR',
                        'devolucaoContainer': 'Paranaguá',
                        'observacoes': "CTNR de 40'",
                        'peso': '14 Ton',
                        'tarifaFrete': '9000.00',
                        'pedagio': '871.20',
                        'adValorem': '0,13%',
                        'gris': '0,07%',
                        'icms': 'Não incluso',
                        'prazoDias': '2',
                    }
                ],
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(created.status_code, 201, created.content)
        body = created.json()
        self.assertEqual(body['valorEstimado'], '9871.20')
        self.assertEqual(len(body['linhas']), 1)
        self.assertEqual(body['linhas'][0]['totalEstimado'], '9871.20')
        self.assertGreaterEqual(len(body['condicoes']), 19)
        self.assertEqual(body['dataVencimento'], '2026-03-30')

    def test_proposta_usa_razao_social_mesmo_com_nome_fantasia(self):
        self._auth(self.admin)
        cliente = self.api.post('/api/comercial/clientes/', PAYLOAD, format='json', **HEADERS)
        self.assertEqual(cliente.status_code, 201, cliente.content)
        created = self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'armazenagem',
                'clienteId': cliente.json()['id'],
                'clienteNome': 'Empresa Teste',
                'status': 'enviada',
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(created.status_code, 201, created.content)
        self.assertEqual(created.json()['clienteNome'], 'EMPRESA TESTE LTDA')
        listed = self.api.get('/api/comercial/propostas/', **HEADERS)
        self.assertEqual(listed.json()['results'][0]['clienteNome'], 'EMPRESA TESTE LTDA')

    def test_operador_sem_funcao_nao_cria_proposta(self):
        self._auth(self.leitura)
        response = self.api.post(
            '/api/comercial/propostas/',
            {'tipo': 'frete', 'titulo': 'Sem permissão'},
            format='json',
            **HEADERS,
        )
        self.assertEqual(response.status_code, 403)

    def test_enviar_email_exige_google_vinculado(self):
        self._auth(self.admin)
        cliente = self.api.post(
            '/api/comercial/clientes/',
            {**PAYLOAD, 'email': 'compras@empresa.com'},
            format='json',
            **HEADERS,
        )
        proposta = self.api.post(
            '/api/comercial/propostas/',
            {'tipo': 'armazenagem', 'clienteId': cliente.json()['id'], 'status': 'rascunho'},
            format='json',
            **HEADERS,
        )
        response = self.api.post(
            f'/api/comercial/propostas/{proposta.json()["id"]}/enviar-email/',
            {},
            format='json',
            **HEADERS,
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn('Google', response.json()['detail'])

    @patch('apps.comercial.proposta_email_service.send_gmail_as_user')
    def test_enviar_email_pelo_gmail_do_usuario(self, mock_send):
        self.admin.google_email = 'miguel.ribeiro@transcamila.com.br'
        self.admin.save(update_fields=['google_email'])
        self._auth(self.admin)
        cliente = self.api.post(
            '/api/comercial/clientes/',
            {**PAYLOAD, 'email': 'compras@empresa.com'},
            format='json',
            **HEADERS,
        )
        proposta = self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'transporte_rodoviario',
                'clienteId': cliente.json()['id'],
                'att': 'Bianca',
                'status': 'rascunho',
                'linhas': [{'origem': 'Ibiporã-PR', 'entrega': 'Rondonópolis', 'veiculo': 'Carreta'}],
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(proposta.status_code, 201, proposta.content)
        proposta_id = proposta.json()['id']
        response = self.api.post(
            f'/api/comercial/propostas/{proposta_id}/enviar-email/',
            {
                'cc': ['diretor@transcamila.com.br', 'miguel.ribeiro@transcamila.com.br'],
                'pdfBase64': PDF_BASE64,
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(response.status_code, 200, response.content)
        mock_send.assert_called_once()
        _user, email_obj = mock_send.call_args.args
        self.assertEqual(email_obj.to, ['compras@empresa.com'])
        self.assertEqual(email_obj.cc, ['diretor@transcamila.com.br'])
        self.assertIn('miguel.ribeiro@transcamila.com.br', email_obj.from_email)
        self.assertTrue(email_obj.attachments)
        self.assertIn('cid:logo_transcamila', email_obj.body)
        pdf_anexo = next(
            (item[0] for item in email_obj.attachments if isinstance(item, tuple) and str(item[0]).endswith('.pdf')),
            '',
        )
        self.assertIn('EMPRESA_TESTE_LTDA', pdf_anexo)
        self.assertIn('001-2026', pdf_anexo)
        self.assertIn('Proposta comercial nº', email_obj.subject)
        atualizada = self.api.get(f'/api/comercial/propostas/{proposta_id}/', **HEADERS)
        self.assertEqual(atualizada.json()['status'], 'enviada')
        self.assertEqual(atualizada.json()['clienteEmail'], 'compras@empresa.com')

    def _pdf_upload(self, name='proposta.pdf'):
        return SimpleUploadedFile(name, base64.b64decode(PDF_BASE64), content_type='application/pdf')

    @patch('apps.comercial.proposta_email_service.send_gmail_as_user')
    def test_enviar_email_lote_frete_e_armazenagem_mesmo_cliente(self, mock_send):
        self.admin.google_email = 'miguel.ribeiro@transcamila.com.br'
        self.admin.save(update_fields=['google_email'])
        self._auth(self.admin)
        cliente = self.api.post(
            '/api/comercial/clientes/',
            {**PAYLOAD, 'email': 'compras@empresa.com'},
            format='json',
            **HEADERS,
        )
        cliente_id = cliente.json()['id']
        frete = self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'transporte_rodoviario',
                'clienteId': cliente_id,
                'status': 'rascunho',
                'linhas': [{'origem': 'Ibiporã-PR', 'entrega': 'Rondonópolis', 'veiculo': 'Carreta'}],
            },
            format='json',
            **HEADERS,
        )
        armazem = self.api.post(
            '/api/comercial/propostas/',
            {'tipo': 'armazenagem', 'clienteId': cliente_id, 'status': 'rascunho'},
            format='json',
            **HEADERS,
        )
        response = self.api.post(
            '/api/comercial/propostas/enviar-email-lote/',
            {
                'ids': [frete.json()['id'], armazem.json()['id']],
                'pdf': [self._pdf_upload('frete.pdf'), self._pdf_upload('armazem.pdf')],
            },
            format='multipart',
            **HEADERS,
        )
        self.assertEqual(response.status_code, 200, response.content)
        mock_send.assert_called_once()
        _user, email_obj = mock_send.call_args.args
        pdfs = [item[0] for item in email_obj.attachments if isinstance(item, tuple) and str(item[0]).endswith('.pdf')]
        self.assertEqual(len(pdfs), 2)
        self.assertIn('Propostas comerciais nº', email_obj.subject)
        self.assertEqual(self.api.get(f'/api/comercial/propostas/{frete.json()["id"]}/', **HEADERS).json()['status'], 'enviada')
        self.assertEqual(self.api.get(f'/api/comercial/propostas/{armazem.json()["id"]}/', **HEADERS).json()['status'], 'enviada')

    def test_enviar_email_lote_rejeita_clientes_diferentes(self):
        self.admin.google_email = 'miguel.ribeiro@transcamila.com.br'
        self.admin.save(update_fields=['google_email'])
        self._auth(self.admin)
        cliente_a = self.api.post(
            '/api/comercial/clientes/',
            {**PAYLOAD, 'email': 'a@empresa.com'},
            format='json',
            **HEADERS,
        )
        cliente_b = self.api.post(
            '/api/comercial/clientes/',
            {**PAYLOAD, 'cnpj': '00.000.000/0002-72', 'razaoSocial': 'Outra Empresa', 'email': 'b@empresa.com'},
            format='json',
            **HEADERS,
        )
        frete = self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'transporte_rodoviario',
                'clienteId': cliente_a.json()['id'],
                'status': 'rascunho',
                'linhas': [{'origem': 'Ibiporã-PR', 'entrega': 'Rondonópolis', 'veiculo': 'Carreta'}],
            },
            format='json',
            **HEADERS,
        )
        armazem = self.api.post(
            '/api/comercial/propostas/',
            {'tipo': 'armazenagem', 'clienteId': cliente_b.json()['id'], 'status': 'rascunho'},
            format='json',
            **HEADERS,
        )
        response = self.api.post(
            '/api/comercial/propostas/enviar-email-lote/',
            {'ids': [frete.json()['id'], armazem.json()['id']]},
            format='multipart',
            **HEADERS,
        )
        self.assertEqual(response.status_code, 400, response.content)
        self.assertIn('mesmo cliente', response.json()['detail'])

    def test_enviar_email_exige_destinatario(self):
        self.admin.google_email = 'miguel.ribeiro@transcamila.com.br'
        self.admin.save(update_fields=['google_email'])
        self._auth(self.admin)
        proposta = self.api.post(
            '/api/comercial/propostas/',
            {'tipo': 'armazenagem', 'titulo': 'Sem cliente', 'status': 'rascunho'},
            format='json',
            **HEADERS,
        )
        response = self.api.post(
            f'/api/comercial/propostas/{proposta.json()["id"]}/enviar-email/',
            {},
            format='json',
            **HEADERS,
        )
        self.assertEqual(response.status_code, 400)

    @patch('apps.comercial.proposta_email_service.send_gmail_as_user')
    def test_enviar_email_nao_coloca_remetente_em_copia(self, mock_send):
        self.admin.google_email = 'adm.ibi@transcamila.com.br'
        self.admin.save(update_fields=['google_email'])
        self._auth(self.admin)
        cliente = self.api.post(
            '/api/comercial/clientes/',
            {**PAYLOAD, 'email': 'compras@empresa.com', 'cnpj': '00.000.000/0002-72'},
            format='json',
            **HEADERS,
        )
        proposta = self.api.post(
            '/api/comercial/propostas/',
            {'tipo': 'armazenagem', 'clienteId': cliente.json()['id'], 'status': 'rascunho'},
            format='json',
            **HEADERS,
        )
        response = self.api.post(
            f'/api/comercial/propostas/{proposta.json()["id"]}/enviar-email/',
            {
                'cc': ['adm.ibi@transcamila.com.br', 'diretor@transcamila.com.br'],
                'pdfBase64': PDF_BASE64,
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(response.status_code, 200, response.content)
        _user, email_obj = mock_send.call_args.args
        self.assertEqual(email_obj.to, ['compras@empresa.com'])
        self.assertEqual(email_obj.cc, ['diretor@transcamila.com.br'])

    def test_aceitar_proposta_promove_potencial_para_cliente(self):
        self._auth(self.admin)
        cliente = self.api.post(
            '/api/comercial/clientes/',
            {**PAYLOAD, 'situacao': 'potencial', 'cnpj': '00.000.000/0003-53'},
            format='json',
            **HEADERS,
        )
        self.assertEqual(cliente.status_code, 201, cliente.content)
        cliente_id = cliente.json()['id']
        self.assertEqual(cliente.json()['situacao'], 'potencial')
        self.assertIsNone(cliente.json()['clienteDesde'])

        proposta = self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'armazenagem',
                'clienteId': cliente_id,
                'status': 'enviada',
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(proposta.status_code, 201, proposta.content)
        atual = self.api.get(f'/api/comercial/clientes/{cliente_id}/', **HEADERS)
        self.assertEqual(atual.json()['situacao'], 'potencial')

        aceita = self.api.patch(
            f'/api/comercial/propostas/{proposta.json()["id"]}/',
            {'status': 'aprovada'},
            format='json',
            **HEADERS,
        )
        self.assertEqual(aceita.status_code, 200, aceita.content)
        promovido = self.api.get(f'/api/comercial/clientes/{cliente_id}/', **HEADERS)
        self.assertEqual(promovido.json()['situacao'], 'cliente')
        self.assertIsNotNone(promovido.json()['clienteDesde'])

        historico = self.api.get(f'/api/comercial/clientes/{cliente_id}/historico/', **HEADERS)
        self.assertEqual(historico.status_code, 200, historico.content)
        body = historico.json()
        self.assertEqual(body['situacao'], 'cliente')
        self.assertIsNotNone(body['clienteDesde'])
        self.assertEqual(len(body['propostasAceitas']), 1)
        self.assertEqual(body['propostasAceitas'][0]['id'], proposta.json()['id'])
        self.assertEqual(body['totalPropostas'], 1)
        self.assertEqual(body['propostasAceitasCount'], 1)
        self.assertEqual(body['propostasRecusadasCount'], 0)
        self.assertEqual(body['indiceAceitacao'], 100.0)

    def test_catalogo_tabela_frete_e_generalidades(self):
        self._auth(self.admin)
        cliente = self.api.post('/api/comercial/clientes/', PAYLOAD, format='json', **HEADERS)
        self.assertEqual(cliente.status_code, 201, cliente.content)
        cliente_id = cliente.json()['id']
        outro = self.api.post(
            '/api/comercial/clientes/',
            {**PAYLOAD, 'cnpj': '60.701.190/0001-04', 'razaoSocial': 'Outro Cliente SA'},
            format='json',
            **HEADERS,
        )
        self.assertEqual(outro.status_code, 201, outro.content)
        outro_id = outro.json()['id']

        missing = self.api.get('/api/comercial/generalidades/', **HEADERS)
        self.assertEqual(missing.status_code, 400)

        padrao_geral = self.api.get('/api/comercial/generalidades/?tipo=frete', **HEADERS)
        self.assertEqual(padrao_geral.status_code, 200, padrao_geral.content)
        self.assertEqual(padrao_geral.json()['origem'], 'padrao')
        self.assertIsNone(padrao_geral.json()['clienteId'])
        self.assertEqual(padrao_geral.json()['items'][0]['rotulo'], 'Capacidade dos veículos - Carreta')

        seeded = self.api.get(
            f'/api/comercial/generalidades/?cliente={cliente_id}&tipo=distribuicao',
            **HEADERS,
        )
        self.assertEqual(seeded.status_code, 200)
        self.assertEqual(seeded.json()['origem'], 'padrao')
        items = seeded.json()['items']
        self.assertGreaterEqual(len(items), 16)
        self.assertEqual(items[0]['rotulo'], 'Limite de embarque')

        padrao_frete = self.api.get(
            f'/api/comercial/generalidades/?cliente={cliente_id}&tipo=frete',
            **HEADERS,
        )
        self.assertEqual(padrao_frete.status_code, 200)
        self.assertEqual(padrao_frete.json()['origem'], 'padrao')
        self.assertEqual(padrao_frete.json()['items'][0]['rotulo'], 'Capacidade dos veículos - Carreta')

        saved = self.api.put(
            '/api/comercial/generalidades/',
            {
                'clienteId': cliente_id,
                'tipo': 'distribuicao',
                'items': [{'rotulo': 'Pedágios', 'valor': 'Conforme Legislação'}],
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(saved.status_code, 200, saved.content)
        self.assertEqual(saved.json()['origem'], 'cliente')
        self.assertEqual(len(saved.json()['items']), 1)

        stored = self.api.get(
            f'/api/comercial/generalidades/?cliente={cliente_id}&tipo=distribuicao',
            **HEADERS,
        )
        self.assertEqual(stored.json()['items'][0]['rotulo'], 'Pedágios')

        outro_tipo = self.api.get(
            f'/api/comercial/generalidades/?cliente={cliente_id}&tipo=armazenagem',
            **HEADERS,
        )
        self.assertEqual(outro_tipo.json()['origem'], 'padrao')
        self.assertGreaterEqual(len(outro_tipo.json()['items']), 15)
        self.assertEqual(outro_tipo.json()['items'][0]['rotulo'], '(1)')

        frete = self.api.put(
            '/api/comercial/generalidades/',
            {
                'clienteId': cliente_id,
                'tipo': 'frete',
                'items': [{'rotulo': 'Pedágio', 'valor': 'Por trecho'}],
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(frete.status_code, 200, frete.content)
        self.assertEqual(len(frete.json()['items']), 1)

        stored_dist = self.api.get(
            f'/api/comercial/generalidades/?cliente={cliente_id}&tipo=distribuicao',
            **HEADERS,
        )
        self.assertEqual(stored_dist.json()['items'][0]['rotulo'], 'Pedágios')

        outro_cliente = self.api.get(
            f'/api/comercial/generalidades/?cliente={outro_id}&tipo=distribuicao',
            **HEADERS,
        )
        self.assertEqual(outro_cliente.json()['origem'], 'padrao')

        padrao_editado = self.api.put(
            '/api/comercial/generalidades/',
            {
                'tipo': 'armazenagem',
                'aplicar': 'todos',
                'items': [{'rotulo': '(1)', 'valor': 'Padrão geral atualizado'}],
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(padrao_editado.status_code, 200, padrao_editado.content)
        self.assertEqual(padrao_editado.json()['origem'], 'padrao')
        self.assertEqual(padrao_editado.json()['items'][0]['valor'], 'Padrão geral atualizado')
        herdado = self.api.get(
            f'/api/comercial/generalidades/?cliente={outro_id}&tipo=armazenagem',
            **HEADERS,
        )
        self.assertEqual(herdado.json()['origem'], 'padrao')
        self.assertEqual(herdado.json()['items'][0]['valor'], 'Padrão geral atualizado')

        somente_novos = self.api.put(
            '/api/comercial/generalidades/',
            {
                'tipo': 'armazenagem',
                'aplicar': 'novos',
                'items': [{'rotulo': '(1)', 'valor': 'Só para clientes novos'}],
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(somente_novos.status_code, 200, somente_novos.content)
        self.assertEqual(somente_novos.json()['items'][0]['valor'], 'Só para clientes novos')
        cliente_atual = self.api.get(
            f'/api/comercial/generalidades/?cliente={outro_id}&tipo=armazenagem',
            **HEADERS,
        )
        self.assertEqual(cliente_atual.json()['origem'], 'cliente')
        self.assertEqual(cliente_atual.json()['items'][0]['valor'], 'Padrão geral atualizado')

        created = self.api.post(
            '/api/comercial/tabela-frete/',
            {
                'nome': 'CropChem Transferência',
                'tipo': 'transferencia',
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(created.status_code, 201, created.content)
        tabela_id = created.json()['id']
        linha = self.api.post(
            '/api/comercial/tabela-frete-linhas/',
            {
                'tabelaId': tabela_id,
                'origem': 'Ibiporã-PR',
                'entrega': 'Rondonópolis-MT',
                'veiculo': 'Carreta Graneleira',
                'tarifaFrete': '8500.00',
                'prazoDias': '3 dias úteis',
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(linha.status_code, 201, linha.content)
        listed = self.api.get('/api/comercial/tabela-frete/?search=CropChem', **HEADERS)
        self.assertEqual(listed.status_code, 200)
        self.assertEqual(listed.json()['count'], 1)

        dist = self.api.post(
            '/api/comercial/tabela-frete/',
            {
                'nome': 'CropChem Distribuição',
                'tipo': 'distribuicao',
                'config': {'kmInicio': 0, 'kmFim': 100, 'tarifaBase': '240.34'},
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(dist.status_code, 201, dist.content)
        faixas = dist.json()['faixas']
        self.assertGreaterEqual(len(faixas), 2)
        self.assertEqual(faixas[0]['kmDe'], 0)
        self.assertEqual(faixas[0]['kmAte'], 50)
        self.assertEqual(faixas[0]['tarifas'][0]['valor'], '240.34')
        self.assertEqual(faixas[1]['kmDe'], 51)
        self.assertEqual(faixas[1]['kmAte'], 100)

        dist_passos = self.api.post(
            '/api/comercial/tabela-frete/',
            {
                'nome': 'CropChem Passos',
                'tipo': 'distribuicao',
                'config': {'kmInicio': 0, 'kmFim': 300, 'tarifaBase': '240.34'},
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(dist_passos.status_code, 201, dist_passos.content)
        faixas_passos = dist_passos.json()['faixas']
        self.assertEqual([(item['kmDe'], item['kmAte']) for item in faixas_passos], [
            (0, 50),
            (51, 100),
            (101, 150),
            (151, 200),
            (201, 300),
        ])
        self.assertEqual(faixas_passos[4]['tarifas'][0]['valor'], '730.83')

        vinculada = self.api.post(
            '/api/comercial/tabela-frete/',
            {
                'nome': 'Tabela do cliente',
                'tipo': 'distribuicao',
                'clienteIds': [cliente_id],
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(vinculada.status_code, 201, vinculada.content)
        self.assertEqual(vinculada.json()['clienteIds'], [cliente.json()['id']])
        self.assertEqual(vinculada.json()['criadoPorNome'], 'Admin Comercial')
        self.assertEqual(vinculada.json()['atualizadoPorNome'], 'Admin Comercial')
        por_cliente = self.api.get(
            f'/api/comercial/tabela-frete/?tipo=distribuicao&cliente={cliente.json()["id"]}',
            **HEADERS,
        )
        self.assertEqual(por_cliente.status_code, 200)
        self.assertEqual(por_cliente.json()['count'], 1)
        self.assertEqual(por_cliente.json()['results'][0]['id'], vinculada.json()['id'])

        cliente_b = self.api.post(
            '/api/comercial/clientes/',
            {**PAYLOAD, 'cnpj': '00000000000272', 'razaoSocial': 'EMPRESA B LTDA'},
            format='json',
            **HEADERS,
        )
        self.assertEqual(cliente_b.status_code, 201, cliente_b.content)
        multi = self.api.patch(
            f'/api/comercial/tabela-frete/{vinculada.json()["id"]}/',
            {'clienteIds': [cliente.json()['id'], cliente_b.json()['id']]},
            format='json',
            **HEADERS,
        )
        self.assertEqual(multi.status_code, 200, multi.content)
        self.assertEqual(len(multi.json()['clienteIds']), 2)
        self.assertEqual(len(multi.json()['clientesNomes']), 2)

        ocupado = self.api.post(
            '/api/comercial/tabela-frete/',
            {
                'nome': 'Outra tabela',
                'tipo': 'distribuicao',
                'clienteIds': [cliente.json()['id']],
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(ocupado.status_code, 400, ocupado.content)
        self.assertIn('vinculado', str(ocupado.json()).lower())

        published = self.api.post(
            f'/api/comercial/tabela-frete/{vinculada.json()["id"]}/publicar/',
            **HEADERS,
        )
        self.assertEqual(published.status_code, 200, published.content)
        proposta = self.api.post(
            '/api/comercial/propostas/',
            {
                'tipo': 'transporte_rodoviario',
                'clienteId': cliente_id,
                'incluiDistribuicao': True,
                'status': 'rascunho',
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(proposta.status_code, 201, proposta.content)
        self.assertEqual(proposta.json()['condicoes'], [
            {'rotulo': 'Pedágios', 'valor': 'Conforme Legislação', 'tipo': 'distribuicao'},
        ])

        restaurado = self.api.delete(
            f'/api/comercial/generalidades/?cliente={cliente_id}&tipo=distribuicao',
            **HEADERS,
        )
        self.assertEqual(restaurado.status_code, 200, restaurado.content)
        self.assertEqual(restaurado.json()['origem'], 'padrao')

        self._auth(self.leitura)
        denied = self.api.post(
            '/api/comercial/tabela-frete/',
            {'nome': 'Sem permissão', 'tipo': 'transferencia'},
            format='json',
            **HEADERS,
        )
        self.assertEqual(denied.status_code, 403)

    def test_ccab_distribuicao_calculo(self):
        from apps.comercial.tabela_distribuicao import (
            gerar_faixas_distribuicao,
            preset_config_ccab,
            simular_cotacao_distribuicao,
        )

        faixas = gerar_faixas_distribuicao(preset_config_ccab())
        primeira = faixas[0]
        self.assertEqual(primeira['kmDe'], 0)
        self.assertEqual(primeira['kmAte'], 50)
        self.assertEqual(primeira['tarifas'][0]['valor'], '240.34')
        self.assertEqual(primeira['tarifas'][6]['valor'], '1201.70')
        self.assertEqual(primeira['freteMinimo'], '119.93')
        self.assertEqual(primeira['pedagioTon'], '17.00')

        faixa_201 = next(item for item in faixas if item['kmDe'] == 201)
        self.assertEqual(faixa_201['kmAte'], 300)
        self.assertEqual(faixa_201['tarifas'][0]['valor'], '730.83')
        self.assertEqual(faixa_201['freteMinimo'], '364.68')
        self.assertEqual(faixa_201['tarifas'][7]['valor'], '5481.22')
        self.assertEqual(faixa_201['tarifas'][8]['valor'], '6029.34')

        self._auth(self.admin)
        preset = self.api.get('/api/comercial/tabela-frete/preset-ccab/', **HEADERS)
        self.assertEqual(preset.status_code, 200, preset.content)
        body = preset.json()
        self.assertEqual(body['config']['modoTarifa'], 'linear_km_ate')
        self.assertEqual(body['faixas'][0]['tarifas'][0]['valor'], '240.34')

        sim = simular_cotacao_distribuicao(
            preset_config_ccab(),
            1350,
            '1506.13',
            'fracionado',
            '500000',
        )
        self.assertEqual(sim['bandaPeso']['key'], 'de1000')
        self.assertEqual(sim['freteBase'], '3107.99')
        self.assertTrue(sim['grisAdvUnificado'])
        self.assertEqual(sim['grisAdv'], '750.00')
        self.assertEqual(sim['gris'], '0.00')
        self.assertEqual(sim['adv'], '0.00')
        self.assertEqual(sim['subtotal'], '3857.99')
        self.assertEqual(sim['valorPorKm'], '2.30')
        self.assertEqual(sim['total'], '3906.26')
        self.assertEqual(sim['pesoKg'], 1506.13)

        from apps.comercial.simulacao_icms import anexar_icms_simulacao

        sim_icms = anexar_icms_simulacao(
            simular_cotacao_distribuicao(
                preset_config_ccab(),
                250,
                850,
                'fracionado',
                None,
            ),
            'SP',
            'PR',
        )
        self.assertIsNotNone(sim_icms['icms'])
        self.assertEqual(sim_icms['icms']['ufOrigem'], 'SP')
        self.assertEqual(sim_icms['icms']['ufDestino'], 'PR')
        self.assertEqual(sim_icms['icms']['tipo'], 'interestadual')
        self.assertEqual(sim_icms['icms']['aliquotaPercent'], 12)

        sim_interno = anexar_icms_simulacao(
            simular_cotacao_distribuicao(
                preset_config_ccab(),
                250,
                850,
                'fracionado',
                None,
            ),
            'SP',
            'SP',
        )
        self.assertEqual(sim_interno['icms']['tipo'], 'interno')
        self.assertEqual(sim_interno['icms']['aliquotaPercent'], 18)

    def test_albaugh_distribuicao_calculo(self):
        from apps.comercial.tabela_distribuicao import (
            gerar_faixas_distribuicao,
            preset_config_albaugh,
            simular_cotacao_distribuicao,
        )

        config = preset_config_albaugh()
        faixas = gerar_faixas_distribuicao(config)
        self.assertEqual(len(faixas), 30)

        faixa_1400 = next(item for item in faixas if item['kmDe'] == 1301)
        self.assertEqual(faixa_1400['kmAte'], 1400)
        carreta27 = next(t for t in faixa_1400['tarifas'] if t['key'] == 'fechada_carreta27')
        self.assertEqual(carreta27['valor'], '16877.73')
        self.assertEqual(faixa_1400['pedagioTon'], '51.70')

        sim_fechado = simular_cotacao_distribuicao(config, 1400, 5000, 'fechado', '300000')
        self.assertEqual(sim_fechado['freteBase'], '16877.73')
        self.assertEqual(sim_fechado['pedagio'], '258.50')
        self.assertEqual(sim_fechado['adv'], '390.00')
        self.assertFalse(sim_fechado['grisAdvUnificado'])
        self.assertEqual(sim_fechado['bandaPeso']['key'], 'fechada_carreta27')

        sim_fracionado = simular_cotacao_distribuicao(config, 1400, 3000, 'fracionado')
        self.assertEqual(sim_fracionado['freteBase'], '5063.31')
        self.assertEqual(sim_fracionado['bandaPeso']['key'], 'fracionado_ate5t')

        self._auth(self.admin)
        preset = self.api.get('/api/comercial/tabela-frete/preset-albaugh/', **HEADERS)
        self.assertEqual(preset.status_code, 200, preset.content)
        self.assertEqual(len(preset.json()['faixas']), 30)

    def test_colunas_extras_personalizadas(self):
        from apps.comercial.tabela_distribuicao import (
            gerar_faixas_distribuicao,
            preset_config_ccab,
            simular_cotacao_distribuicao,
        )

        config = preset_config_ccab()
        config['colunasExtras'] = [
            {'rotulo': 'Tx emissão CTe', 'calculo': 'fixo', 'valor': '25'},
            {'rotulo': 'Taxa NF', 'calculo': 'percentual_nf', 'valor': '0.10'},
        ]
        faixas = gerar_faixas_distribuicao(config)
        extras = {item['key']: item for item in faixas[0]['extras']}
        self.assertEqual(extras['tx_emissao_cte']['valor'], '25.00')
        self.assertEqual(extras['tx_emissao_cte']['formato'], 'moeda')
        self.assertEqual(extras['taxa_nf']['valor'], '0.10')
        self.assertEqual(extras['taxa_nf']['formato'], 'percentual')

        sim = simular_cotacao_distribuicao(config, 1350, '1506.13', 'fracionado', '500000')
        self.assertEqual(sim['extras'][0]['valor'], '25.00')
        self.assertEqual(sim['extras'][1]['valor'], '500.00')
        self.assertEqual(sim['total'], '4431.26')

        sim_sem_nf = simular_cotacao_distribuicao(config, 1350, '1506.13', 'fracionado')
        self.assertEqual(sim_sem_nf['extras'][1]['valor'], '0.00')
        self.assertEqual(sim_sem_nf['total'], '3181.26')

    def test_gris_adv_unificado_ou_separado(self):
        from apps.comercial.tabela_distribuicao import (
            default_config_distribuicao,
            merge_config,
            preset_config_ccab,
            simular_cotacao_distribuicao,
        )

        legado = merge_config({'grisAdvPercent': '0.0015', 'grisPercent': '0.05', 'advPercent': '0.09'})
        self.assertTrue(legado['grisAdvUnificado'])
        self.assertEqual(legado['grisPercent'], '0.0015')
        self.assertEqual(legado['advPercent'], '0.0015')

        separado = merge_config({
            **default_config_distribuicao(),
            'modoTarifa': 'linear_km_ate',
            'grisAdvUnificado': False,
            'grisPercent': '0.05',
            'advPercent': '0.09',
            'grisAdvPercent': '0.0015',
        })
        self.assertFalse(separado['grisAdvUnificado'])
        self.assertEqual(separado['grisAdvPercent'], '')
        self.assertEqual(separado['grisPercent'], '0.05')
        self.assertEqual(separado['advPercent'], '0.09')

        sim_separado = simular_cotacao_distribuicao(separado, 100, 1000, 'fracionado', '1000')
        self.assertFalse(sim_separado['grisAdvUnificado'])
        self.assertEqual(sim_separado['gris'], '50.00')
        self.assertEqual(sim_separado['adv'], '90.00')

        ccab = merge_config(preset_config_ccab())
        self.assertTrue(ccab['grisAdvUnificado'])
        sim_ccab = simular_cotacao_distribuicao(ccab, 1350, '1506.13', 'fracionado', '500000')
        self.assertTrue(sim_ccab['grisAdvUnificado'])
        self.assertEqual(sim_ccab['grisAdv'], '750.00')
        self.assertEqual(sim_ccab['gris'], '0.00')
        self.assertEqual(sim_ccab['adv'], '0.00')

    def test_contrato_tarifario_publicar_e_simular(self):
        from apps.comercial.tabela_distribuicao import preset_config_ccab

        self._auth(self.admin)
        created = self.api.post(
            '/api/comercial/tabela-frete/',
            {
                'nome': 'CCAB Contrato',
                'codigo': 'CCAB-2024',
                'revisao': 61,
                'tipo': 'distribuicao',
                'config': preset_config_ccab(),
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(created.status_code, 201, created.content)
        tabela_id = created.json()['id']
        self.assertEqual(created.json()['status'], 'rascunho')

        published = self.api.post(f'/api/comercial/tabela-frete/{tabela_id}/publicar/', **HEADERS)
        self.assertEqual(published.status_code, 200, published.content)
        self.assertEqual(published.json()['status'], 'publicada')

        simulacao = self.api.post(
            f'/api/comercial/tabela-frete/{tabela_id}/simular/',
            {
                'km': 1350,
                'pesoKg': 1506.13,
                'modalidade': 'fracionado',
                'valorNf': '500000',
                'ufOrigem': 'AC',
                'ufDestino': 'SC',
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(simulacao.status_code, 200, simulacao.content)
        body = simulacao.json()
        self.assertEqual(body['total'], '3906.26')
        self.assertIsNotNone(body.get('icms'))
        self.assertEqual(body['icms']['ufOrigem'], 'AC')
        self.assertEqual(body['icms']['ufDestino'], 'SC')
        self.assertEqual(body['icms']['aliquotaPercent'], 12)
        self.assertEqual(body['icms']['valor'], '532.67')
        self.assertEqual(body['icms']['totalComIcms'], '4438.93')
        self.assertTrue(body['icms']['incluiPedagioNaBase'])

        simulacao_pr = self.api.post(
            f'/api/comercial/tabela-frete/{tabela_id}/simular/',
            {
                'km': 1350,
                'pesoKg': 1506.13,
                'modalidade': 'fracionado',
                'valorNf': '500000',
                'ufOrigem': 'PR',
                'ufDestino': 'MG',
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(simulacao_pr.status_code, 200, simulacao_pr.content)
        body_pr = simulacao_pr.json()
        self.assertEqual(body_pr['subtotal'], '3857.99')
        self.assertEqual(body_pr['pedagio'], '48.27')
        self.assertEqual(body_pr['icms']['aliquotaPercent'], 12)
        self.assertFalse(body_pr['icms']['incluiPedagioNaBase'])
        self.assertEqual(body_pr['icms']['valor'], '526.09')
        self.assertEqual(body_pr['icms']['totalComIcms'], '4432.35')

        simulacao_sem_icms = self.api.post(
            f'/api/comercial/tabela-frete/{tabela_id}/simular/',
            {'km': 250, 'pesoKg': 850, 'modalidade': 'fracionado', 'valorNf': '10000'},
            format='json',
            **HEADERS,
        )
        self.assertEqual(simulacao_sem_icms.status_code, 200, simulacao_sem_icms.content)
        self.assertEqual(simulacao_sem_icms.json()['bandaPeso']['key'], 'de500')
        self.assertIn('total', simulacao_sem_icms.json())
        self.assertIsNone(simulacao_sem_icms.json().get('icms'))

        revisao = self.api.post(f'/api/comercial/tabela-frete/{tabela_id}/nova-revisao/', **HEADERS)
        self.assertEqual(revisao.status_code, 201, revisao.content)
        self.assertEqual(revisao.json()['revisao'], 62)
        self.assertEqual(revisao.json()['status'], 'rascunho')

        historico = self.api.get(f'/api/comercial/tabela-frete/{tabela_id}/historico-revisoes/', **HEADERS)
        self.assertEqual(historico.status_code, 200, historico.content)
        itens = historico.json()['revisoes']
        self.assertEqual(len(itens), 2)
        rev61 = next(item for item in itens if item['revisao'] == 61)
        rev62 = next(item for item in itens if item['revisao'] == 62)
        self.assertTrue(rev61['atual'])
        self.assertFalse(rev62['atual'])
        self.assertEqual(rev61['status'], 'arquivada')
        self.assertEqual(rev62['usuarioNome'], 'Admin Comercial')

        nova_id = revisao.json()['id']
        historico_nova = self.api.get(f'/api/comercial/tabela-frete/{nova_id}/historico-revisoes/', **HEADERS)
        self.assertEqual(historico_nova.status_code, 200, historico_nova.content)
        rev62_atual = next(item for item in historico_nova.json()['revisoes'] if item['revisao'] == 62)
        self.assertTrue(rev62_atual['atual'])

        listed = self.api.get('/api/comercial/tabela-frete/?search=CCAB-2024', **HEADERS)
        self.assertEqual(listed.status_code, 200, listed.content)
        ccab_rows = [item for item in listed.json()['results'] if item.get('codigo') == 'CCAB-2024']
        self.assertEqual(len(ccab_rows), 1)
        self.assertEqual(ccab_rows[0]['revisao'], 62)
        self.assertEqual(ccab_rows[0]['id'], nova_id)

        old_detail = self.api.get(f'/api/comercial/tabela-frete/{tabela_id}/', **HEADERS)
        self.assertEqual(old_detail.status_code, 200, old_detail.content)
        self.assertEqual(old_detail.json()['revisao'], 61)
        self.assertEqual(old_detail.json()['status'], 'arquivada')

    def test_listagem_uma_linha_por_codigo(self):
        self._auth(self.admin)
        cliente = self.api.post('/api/comercial/clientes/', PAYLOAD, format='json', **HEADERS)
        self.assertEqual(cliente.status_code, 201, cliente.content)
        created = self.api.post(
            '/api/comercial/tabela-frete/',
            {
                'nome': 'CCAB Distribuição',
                'codigo': 'CCAB-2024',
                'revisao': 70,
                'tipo': 'distribuicao',
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(created.status_code, 201, created.content)
        tabela_id = created.json()['id']
        published = self.api.post(f'/api/comercial/tabela-frete/{tabela_id}/publicar/', **HEADERS)
        self.assertEqual(published.status_code, 200, published.content)
        revisao = self.api.post(f'/api/comercial/tabela-frete/{tabela_id}/nova-revisao/', **HEADERS)
        self.assertEqual(revisao.status_code, 201, revisao.content)
        nova_id = revisao.json()['id']
        vinculados = self.api.patch(
            f'/api/comercial/tabela-frete/{nova_id}/',
            {'clienteIds': [cliente.json()['id']]},
            format='json',
            **HEADERS,
        )
        self.assertEqual(vinculados.status_code, 200, vinculados.content)

        listed = self.api.get('/api/comercial/tabela-frete/?tipo=distribuicao&search=CCAB-2024', **HEADERS)
        self.assertEqual(listed.status_code, 200, listed.content)
        linhas = [item for item in listed.json()['results'] if item.get('codigo') == 'CCAB-2024']
        self.assertEqual(len(linhas), 1)
        self.assertEqual(linhas[0]['id'], nova_id)
        self.assertEqual(linhas[0]['revisao'], 71)

    def test_excluir_tabela_frete_remove_todas_revisoes(self):
        self._auth(self.admin)
        created = self.api.post(
            '/api/comercial/tabela-frete/',
            {
                'nome': 'CCAB Distribuição',
                'codigo': 'CCAB-EXCL-2024',
                'revisao': 70,
                'tipo': 'distribuicao',
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(created.status_code, 201, created.content)
        tabela_id = created.json()['id']
        published = self.api.post(f'/api/comercial/tabela-frete/{tabela_id}/publicar/', **HEADERS)
        self.assertEqual(published.status_code, 200, published.content)
        revisao = self.api.post(f'/api/comercial/tabela-frete/{tabela_id}/nova-revisao/', **HEADERS)
        self.assertEqual(revisao.status_code, 201, revisao.content)
        nova_id = revisao.json()['id']

        deleted = self.api.delete(f'/api/comercial/tabela-frete/{nova_id}/', **HEADERS)
        self.assertEqual(deleted.status_code, 204, deleted.content)
        listed = self.api.get('/api/comercial/tabela-frete/?search=CCAB-EXCL-2024', **HEADERS)
        self.assertEqual(listed.status_code, 200, listed.content)
        self.assertEqual(
            [item for item in listed.json()['results'] if item.get('codigo') == 'CCAB-EXCL-2024'],
            [],
        )
        self.assertEqual(self.api.get(f'/api/comercial/tabela-frete/{tabela_id}/', **HEADERS).status_code, 404)
        self.assertEqual(self.api.get(f'/api/comercial/tabela-frete/{nova_id}/', **HEADERS).status_code, 404)

    def test_exportar_tabela_frete(self):
        from io import BytesIO

        import openpyxl
        from apps.comercial.tabela_distribuicao import preset_config_ccab

        self._auth(self.admin)
        created = self.api.post(
            '/api/comercial/tabela-frete/',
            {
                'nome': 'CCAB Distribuição',
                'codigo': 'CCAB-2024',
                'revisao': 65,
                'tipo': 'distribuicao',
                'config': preset_config_ccab(),
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(created.status_code, 201, created.content)
        tabela_id = created.json()['id']

        exported = self.api.get(f'/api/comercial/tabela-frete/{tabela_id}/exportar/', **HEADERS)
        self.assertEqual(exported.status_code, 200, exported.content)
        self.assertIn(
            'spreadsheetml.sheet',
            exported['Content-Type'],
        )
        self.assertIn('CCAB-2024_rev65.xlsx', exported['Content-Disposition'])

        workbook = openpyxl.load_workbook(BytesIO(exported.content))
        self.assertEqual(workbook.sheetnames, ['Identificação', 'Faixas'])
        ident = workbook['Identificação']
        self.assertEqual(ident['A2'].value, 'Nome')
        self.assertEqual(ident['B2'].value, 'CCAB Distribuição')
        faixas = workbook['Faixas']
        self.assertEqual(faixas['A1'].value, 'De')
        self.assertEqual(faixas['C1'].value, 'Frete mín.')
        self.assertIsNotNone(faixas['A2'].value)

        transferencia = self.api.post(
            '/api/comercial/tabela-frete/',
            {'nome': 'Trechos Sul', 'tipo': 'transferencia'},
            format='json',
            **HEADERS,
        )
        self.assertEqual(transferencia.status_code, 201, transferencia.content)
        linha = self.api.post(
            '/api/comercial/tabela-frete-linhas/',
            {
                'tabelaId': transferencia.json()['id'],
                'origem': 'Ibiporã-PR',
                'entrega': 'Rondonópolis-MT',
                'tarifaFrete': '8500.00',
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(linha.status_code, 201, linha.content)
        trechos = self.api.get(
            f'/api/comercial/tabela-frete/{transferencia.json()["id"]}/exportar/',
            **HEADERS,
        )
        self.assertEqual(trechos.status_code, 200, trechos.content)
        wb_trechos = openpyxl.load_workbook(BytesIO(trechos.content))
        self.assertEqual(wb_trechos.sheetnames, ['Identificação', 'Trechos'])
        self.assertEqual(wb_trechos['Trechos']['A2'].value, 'Ibiporã-PR')

        self._auth(self.leitura)
        leitura = self.api.get(f'/api/comercial/tabela-frete/{tabela_id}/exportar/', **HEADERS)
        self.assertEqual(leitura.status_code, 200, leitura.content)

    def test_matriz_icms_ufs_padrao_br(self):
        from apps.comercial.icms_uf import aliquota_interestadual_padrao, aliquotas_por_uf_padrao, normalizar_aliquotas

        self.assertEqual(aliquota_interestadual_padrao('SP', 'BA'), 7)
        self.assertEqual(aliquota_interestadual_padrao('RJ', 'RJ'), 20)
        self.assertEqual(aliquotas_por_uf_padrao()['SP'], 18)

        legado = {'SP': {'SP': 18, 'BA': 7}, 'RJ': {'RJ': 20}}
        self.assertEqual(normalizar_aliquotas(legado)['SP'], 18)

        self._auth(self.admin)
        resp = self.api.get('/api/comercial/icms-ufs/', **HEADERS)
        self.assertEqual(resp.status_code, 200, resp.content)
        body = resp.json()
        self.assertEqual(len(body['ufs']), 27)
        self.assertEqual(len(body['regioes']), 5)
        self.assertEqual(body['aliquotas']['RJ'], 20)

        self._auth(self.leitura)
        denied = self.api.put('/api/comercial/icms-ufs/', {'aliquotas': body['aliquotas']}, format='json', **HEADERS)
        self.assertEqual(denied.status_code, 403)

        self._auth(self.admin)
        save = self.api.put(
            '/api/comercial/icms-ufs/',
            {'aliquotas': {**body['aliquotas'], 'SP': 19}},
            format='json',
            **HEADERS,
        )
        self.assertEqual(save.status_code, 200, save.content)
        self.assertEqual(save.json()['aliquotas']['SP'], 19)

        restore = self.api.post('/api/comercial/icms-ufs/', {'acao': 'restaurar-padrao'}, format='json', **HEADERS)
        self.assertEqual(restore.status_code, 200, restore.content)
        self.assertEqual(restore.json()['aliquotas']['SP'], 18)

    def test_rota_distancia_cidades(self):
        from apps.comercial import distancia_rota as rota_mod

        def fake_get_json(url, timeout=12):
            if 'nominatim' in url:
                if 'Curitiba' in url:
                    return [{'lat': '-25.4284', 'lon': '-49.2733', 'display_name': 'Curitiba, PR, Brasil'}]
                return [{'lat': '-23.5505', 'lon': '-46.6333', 'display_name': 'São Paulo, SP, Brasil'}]
            return {
                'code': 'Ok',
                'routes': [{'distance': 408_500}],
            }

        with patch.object(rota_mod, 'usa_google_maps', return_value=False), patch.object(rota_mod, '_get_json', side_effect=fake_get_json):
            self._auth(self.admin)
            resp = self.api.post(
                '/api/comercial/rota-distancia/',
                {
                    'cidadeOrigem': 'São Paulo',
                    'ufOrigem': 'SP',
                    'cidadeDestino': 'Curitiba',
                    'ufDestino': 'PR',
                },
                format='json',
                **HEADERS,
            )
        self.assertEqual(resp.status_code, 200, resp.content)
        body = resp.json()
        self.assertEqual(body['km'], 409)
        self.assertEqual(body['provedor'], 'osm')
        self.assertIn('São Paulo', body['cidadeOrigem'])

    def test_rota_distancia_enderecos(self):
        from apps.comercial import distancia_rota as rota_mod

        def fake_get_json(url, timeout=12):
            if 'nominatim' in url:
                if 'Paulista' in url:
                    return [{'lat': '-23.5614', 'lon': '-46.6559', 'display_name': 'Av. Paulista, São Paulo, SP, Brasil'}]
                return [{'lat': '-25.4284', 'lon': '-49.2733', 'display_name': 'Rua XV de Novembro, Curitiba, PR, Brasil'}]
            return {
                'code': 'Ok',
                'routes': [{'distance': 408_500}],
            }

        with patch.object(rota_mod, 'usa_google_maps', return_value=False), patch.object(rota_mod, '_get_json', side_effect=fake_get_json):
            self._auth(self.admin)
            resp = self.api.post(
                '/api/comercial/rota-distancia/',
                {
                    'modo': 'endereco',
                    'enderecoOrigem': 'Av. Paulista, São Paulo, SP',
                    'enderecoDestino': 'Rua XV de Novembro, Curitiba, PR',
                },
                format='json',
                **HEADERS,
            )
        self.assertEqual(resp.status_code, 200, resp.content)
        body = resp.json()
        self.assertEqual(body['modo'], 'endereco')
        self.assertEqual(body['km'], 409)
        self.assertIn('Paulista', body['enderecoOrigem'])

    def test_buscar_enderecos(self):
        from apps.comercial import distancia_rota as rota_mod

        def fake_get_json(url, timeout=12):
            return [
                {
                    'lat': '-23.5505',
                    'lon': '-46.6333',
                    'display_name': 'São Paulo, SP, Brasil',
                },
            ]

        with patch.object(rota_mod, 'usa_google_maps', return_value=False), patch.object(rota_mod, '_get_json', side_effect=fake_get_json):
            self._auth(self.admin)
            resp = self.api.get('/api/comercial/enderecos/buscar/?q=São Paulo', **HEADERS)
        self.assertEqual(resp.status_code, 200, resp.content)
        body = resp.json()
        self.assertEqual(len(body['results']), 1)
        self.assertIn('São Paulo', body['results'][0]['label'])

    def test_buscar_enderecos_por_empresa(self):
        from apps.comercial import distancia_rota as rota_mod

        def fake_get_json(url, timeout=12):
            if 'place/textsearch' in url:
                return {
                    'status': 'OK',
                    'results': [{
                        'name': 'CCAB Agro',
                        'formatted_address': 'Alameda Santos, 2159 - São Paulo - SP, Brasil',
                        'geometry': {'location': {'lat': -23.5614, 'lng': -46.6559}},
                    }],
                }
            return {'status': 'ZERO_RESULTS', 'results': []}

        with patch.object(rota_mod, 'usa_google_maps', return_value=True), patch.object(rota_mod, '_get_json', side_effect=fake_get_json):
            self._auth(self.admin)
            resp = self.api.get('/api/comercial/enderecos/buscar/?q=ccab', **HEADERS)
        self.assertEqual(resp.status_code, 200, resp.content)
        body = resp.json()
        self.assertEqual(len(body['results']), 1)
        self.assertIn('CCAB Agro', body['results'][0]['label'])

    def test_buscar_cidades_formata_cidade_uf(self):
        from apps.comercial import distancia_rota as rota_mod

        def fake_get_json(url, timeout=12):
            if 'place/autocomplete' in url:
                return {
                    'status': 'OK',
                    'predictions': [{
                        'description': 'Ibiporã, PR, Brazil',
                        'structured_formatting': {
                            'main_text': 'Ibiporã',
                            'secondary_text': 'PR, Brazil',
                        },
                    }],
                }
            return {'status': 'ZERO_RESULTS', 'results': []}

        with patch.object(rota_mod, 'usa_google_maps', return_value=True), patch.object(rota_mod, '_get_json', side_effect=fake_get_json):
            self._auth(self.admin)
            resp = self.api.get('/api/comercial/enderecos/buscar/?q=Ibipora&tipo=cidade', **HEADERS)
        self.assertEqual(resp.status_code, 200, resp.content)
        body = resp.json()
        self.assertEqual(body['results'][0]['label'], 'Ibiporã-PR')

    def test_endereco_reverso(self):
        from apps.comercial import distancia_rota as rota_mod

        def fake_get_json(url, timeout=12):
            return {
                'status': 'OK',
                'results': [{
                    'formatted_address': 'Av. Paulista, São Paulo - SP, Brasil',
                    'geometry': {'location': {'lat': -23.5614, 'lng': -46.6559}},
                    'address_components': [
                        {'short_name': 'SP', 'types': ['administrative_area_level_1']},
                    ],
                }],
            }

        with patch.object(rota_mod, 'usa_google_maps', return_value=True), patch.object(rota_mod, '_get_json', side_effect=fake_get_json):
            self._auth(self.admin)
            resp = self.api.get('/api/comercial/enderecos/reverso/?lat=-23.5614&lon=-46.6559', **HEADERS)
        self.assertEqual(resp.status_code, 200, resp.content)
        body = resp.json()
        self.assertIn('Paulista', body['label'])
        self.assertEqual(body['uf'], 'SP')

    def test_rota_distancia_google_maps(self):
        from apps.comercial import distancia_rota as rota_mod

        def fake_get_json(url, timeout=12):
            if 'geocode' in url:
                if 'Curitiba' in url:
                    return {
                        'status': 'OK',
                        'results': [{
                            'formatted_address': 'Curitiba, PR, Brasil',
                            'geometry': {'location': {'lat': -25.4284, 'lng': -49.2733}},
                        }],
                    }
                return {
                    'status': 'OK',
                    'results': [{
                        'formatted_address': 'São Paulo, SP, Brasil',
                        'geometry': {'location': {'lat': -23.5505, 'lng': -46.6333}},
                    }],
                }
            return {
                'status': 'OK',
                'rows': [{'elements': [{'status': 'OK', 'distance': {'value': 408_500}}]}],
            }

        with patch.object(rota_mod, 'usa_google_maps', return_value=True), patch.object(rota_mod, '_get_json', side_effect=fake_get_json):
            self._auth(self.admin)
            resp = self.api.post(
                '/api/comercial/rota-distancia/',
                {
                    'cidadeOrigem': 'São Paulo',
                    'ufOrigem': 'SP',
                    'cidadeDestino': 'Curitiba',
                    'ufDestino': 'PR',
                },
                format='json',
                **HEADERS,
            )
        self.assertEqual(resp.status_code, 200, resp.content)
        body = resp.json()
        self.assertEqual(body['km'], 409)
        self.assertEqual(body['provedor'], 'google')
        self.assertIn('São Paulo', body['cidadeOrigem'])


class PropostaComercialDraftTests(TestCase):
    def setUp(self):
        self.api = APIClient()
        self.admin = User.objects.create_user(
            username='comercial.draft.admin',
            password='test123',
            name='Admin Draft',
            role_id='1',
            status='ativo',
            environments=['Comercial'],
        )
        self.outro = User.objects.create_user(
            username='comercial.draft.outro',
            password='test123',
            name='Outro Draft',
            role_id='1',
            status='ativo',
            environments=['Comercial'],
        )
        self.leitura = User.objects.create_user(
            username='comercial.draft.view',
            password='test123',
            name='Leitura Draft',
            role_id='2',
            status='ativo',
            environments=['Comercial'],
        )

    def _auth(self, user):
        self.api.force_authenticate(user=user)

    def test_put_salva_e_get_retorna_rascunho(self):
        self._auth(self.admin)
        cliente = self.api.post('/api/comercial/clientes/', PAYLOAD, format='json', **HEADERS)
        self.assertEqual(cliente.status_code, 201, cliente.content)
        response = self.api.put(
            '/api/comercial/propostas/draft/',
            {
                'abaOperacao': 'transferencia',
                'form': {
                    'tipo': 'transporte_rodoviario',
                    'clienteId': cliente.json()['id'],
                    'clienteNome': 'Empresa Teste Ltda',
                    'incluiTransferencia': True,
                    'linhas': [{'origem': 'Ibiporã-PR', 'entrega': 'Curitiba-PR', 'veiculo': 'Carreta'}],
                },
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(response.status_code, 200, response.content)
        self.assertTrue(response.json()['hasDraft'])
        self.assertEqual(response.json()['form']['linhas'][0]['origem'], 'Ibiporã-PR')
        got = self.api.get('/api/comercial/propostas/draft/', **HEADERS)
        self.assertEqual(got.status_code, 200)
        self.assertTrue(got.json()['hasDraft'])
        self.assertEqual(got.json()['form']['clienteId'], cliente.json()['id'])

    def test_draft_isolado_por_usuario(self):
        self._auth(self.admin)
        self.api.put(
            '/api/comercial/propostas/draft/',
            {'form': {'clienteId': '99', 'titulo': 'Meu rascunho'}},
            format='json',
            **HEADERS,
        )
        self._auth(self.outro)
        response = self.api.get('/api/comercial/propostas/draft/', **HEADERS)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.json()['hasDraft'])

    def test_delete_descarta_rascunho(self):
        self._auth(self.admin)
        self.api.put(
            '/api/comercial/propostas/draft/',
            {'form': {'titulo': 'Rascunho'}},
            format='json',
            **HEADERS,
        )
        deleted = self.api.delete('/api/comercial/propostas/draft/', **HEADERS)
        self.assertEqual(deleted.status_code, 204)
        self.assertFalse(PropostaComercialDraft.objects.filter(usuario=self.admin).exists())

    def test_put_vazio_remove_rascunho(self):
        self._auth(self.admin)
        PropostaComercialDraft.objects.create(
            usuario=self.admin,
            payload={'form': {'clienteId': '1', 'titulo': 'X'}},
        )
        response = self.api.put(
            '/api/comercial/propostas/draft/',
            {'form': {'tipo': 'transporte_rodoviario', 'clienteId': '', 'titulo': ''}},
            format='json',
            **HEADERS,
        )
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.json()['hasDraft'])
        self.assertFalse(PropostaComercialDraft.objects.filter(usuario=self.admin).exists())

    def test_operador_sem_funcao_nao_pode_salvar(self):
        self._auth(self.leitura)
        response = self.api.put(
            '/api/comercial/propostas/draft/',
            {'form': {'titulo': 'Bloqueado'}},
            format='json',
            **HEADERS,
        )
        self.assertEqual(response.status_code, 403)


class HomologacaoProdutosComercialTests(TestCase):
    def setUp(self):
        self.api = APIClient()
        self.admin = User.objects.create_user(
            username='homolog.admin',
            password='test123',
            name='Admin Homolog',
            role_id='1',
            status='ativo',
            environments=['Comercial'],
        )
        self.validador = User.objects.create_user(
            username='homolog.val',
            password='test123',
            name='Validador',
            role_id='2',
            status='ativo',
            environments=['Comercial'],
            funcoes={'Comercial': ['validar-clientes']},
        )
        self.produtos = User.objects.create_user(
            username='homolog.prod',
            password='test123',
            name='Cadastro Produtos',
            role_id='2',
            status='ativo',
            environments=['Comercial'],
            funcoes={'Comercial': ['gerenciar-produtos']},
        )

    def _auth(self, user):
        self.api.force_authenticate(user=user)

    def _cliente(self):
        self._auth(self.admin)
        created = self.api.post('/api/comercial/clientes/', PAYLOAD, format='json', **HEADERS)
        self.assertEqual(created.status_code, 201, created.content)
        return created.json()['id']

    def test_vinculo_coloca_cliente_em_pendente_validacao(self):
        cliente_id = self._cliente()
        self._auth(self.produtos)
        created = self.api.post(
            '/api/comercial/produtos/',
            {
                'nome': 'Glifosato 480',
                'numeroOnu': '3082',
                'classeRisco': '9',
                'grupoEmbalagem': 'III',
                'fispq': 'https://exemplo.com/glifosato.pdf',
                'clienteIds': [cliente_id],
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(created.status_code, 201, created.content)
        self.assertEqual(created.json()['numeroOnu'], '3082')
        self.assertEqual(created.json()['classeRisco'], '9')
        self._auth(self.admin)
        cliente = self.api.get(f'/api/comercial/clientes/{cliente_id}/', **HEADERS)
        self.assertEqual(cliente.json()['compatibilidade'], 'pendente_validacao')
        self.assertEqual(cliente.json()['produtosCount'], 1)
        resumo = cliente.json()['homologacaoResumo']
        self.assertEqual(resumo['produtosPerigosos'], 1)
        self.assertTrue(resumo['aptoHomologar'])
        self.assertFalse(resumo['temImpeditivo'])

    def test_homologar_exige_funcao_e_onu_completa(self):
        cliente_id = self._cliente()
        self._auth(self.produtos)
        self.api.post(
            '/api/comercial/produtos/',
            {
                'nome': 'Solvente',
                'numeroOnu': '',
                'classeRisco': '3',
                'clienteIds': [cliente_id],
            },
            format='json',
            **HEADERS,
        )
        self._auth(self.produtos)
        denied = self.api.post(
            f'/api/comercial/clientes/{cliente_id}/homologar/',
            {'decisao': 'homologado'},
            format='json',
            **HEADERS,
        )
        self.assertEqual(denied.status_code, 403)
        self._auth(self.validador)
        blocked = self.api.post(
            f'/api/comercial/clientes/{cliente_id}/homologar/',
            {'decisao': 'homologado'},
            format='json',
            **HEADERS,
        )
        self.assertEqual(blocked.status_code, 400)
        self._auth(self.produtos)
        produto_id = self.api.get('/api/comercial/produtos/', **HEADERS).json()['results'][0]['id']
        self.api.patch(
            f'/api/comercial/produtos/{produto_id}/',
            {'numeroOnu': '1993', 'fispq': 'https://exemplo.com/fds.pdf'},
            format='json',
            **HEADERS,
        )
        self._auth(self.validador)
        ok = self.api.post(
            f'/api/comercial/clientes/{cliente_id}/homologar/',
            {'decisao': 'homologado', 'justificativa': 'Produtos conferidos.'},
            format='json',
            **HEADERS,
        )
        self.assertEqual(ok.status_code, 200, ok.content)
        self.assertEqual(ok.json()['compatibilidade'], 'homologado')

    def test_reprovacao_exige_justificativa(self):
        cliente_id = self._cliente()
        self._auth(self.produtos)
        self.api.post(
            '/api/comercial/produtos/',
            {'nome': 'Ureia', 'clienteIds': [cliente_id]},
            format='json',
            **HEADERS,
        )
        self._auth(self.validador)
        short = self.api.post(
            f'/api/comercial/clientes/{cliente_id}/homologar/',
            {'decisao': 'reprovado', 'justificativa': 'não'},
            format='json',
            **HEADERS,
        )
        self.assertEqual(short.status_code, 400)
        ok = self.api.post(
            f'/api/comercial/clientes/{cliente_id}/homologar/',
            {'decisao': 'reprovado', 'justificativa': 'Produto fora da política de armazenagem da filial.'},
            format='json',
            **HEADERS,
        )
        self.assertEqual(ok.status_code, 200, ok.content)
        self.assertEqual(ok.json()['compatibilidade'], 'reprovado')

    def test_fila_validacao_omite_cliente_sem_produto(self):
        sem_produto = self._cliente()
        self._auth(self.admin)
        outro = self.api.post(
            '/api/comercial/clientes/',
            {**PAYLOAD, 'cnpj': '11.222.333/0001-81', 'razaoSocial': 'Cliente Com Produto Ltda'},
            format='json',
            **HEADERS,
        )
        self.assertEqual(outro.status_code, 201, outro.content)
        com_produto = outro.json()['id']
        self._auth(self.produtos)
        created = self.api.post(
            '/api/comercial/produtos/',
            {'nome': 'Adjuvante', 'clienteIds': [com_produto]},
            format='json',
            **HEADERS,
        )
        self.assertEqual(created.status_code, 201, created.content)
        self._auth(self.validador)
        fila = self.api.get('/api/comercial/clientes/?fila=validacao&page_size=100', **HEADERS)
        self.assertEqual(fila.status_code, 200, fila.content)
        ids = {item['id'] for item in fila.json()['results']}
        self.assertIn(com_produto, ids)
        self.assertNotIn(sem_produto, ids)
        todos = self.api.get('/api/comercial/clientes/?com_produtos=1&page_size=100', **HEADERS)
        self.assertEqual(todos.status_code, 200, todos.content)
        ids_produtos = {item['id'] for item in todos.json()['results']}
        self.assertIn(com_produto, ids_produtos)
        self.assertNotIn(sem_produto, ids_produtos)

    def test_validacao_resumo_e_filtro_impeditivo(self):
        cliente_id = self._cliente()
        self._auth(self.produtos)
        self.api.post(
            '/api/comercial/produtos/',
            {
                'nome': 'Solvente fila',
                'numeroOnu': '',
                'classeRisco': '3',
                'clienteIds': [cliente_id],
            },
            format='json',
            **HEADERS,
        )
        self._auth(self.validador)
        resumo = self.api.get('/api/comercial/clientes/validacao-resumo/', **HEADERS)
        self.assertEqual(resumo.status_code, 200, resumo.content)
        body = resumo.json()
        self.assertGreaterEqual(body['pendentes'], 1)
        self.assertGreaterEqual(body['comImpeditivo'], 1)
        fila = self.api.get('/api/comercial/clientes/?fila=validacao&pendencia=impeditivo&page_size=100', **HEADERS)
        self.assertEqual(fila.status_code, 200, fila.content)
        ids = {item['id'] for item in fila.json()['results']}
        self.assertIn(cliente_id, ids)
        item = next(row for row in fila.json()['results'] if row['id'] == cliente_id)
        self.assertTrue(item['homologacaoResumo']['temImpeditivo'])
        self.assertEqual(item['produtos'][0]['conformidade']['status'], 'bloqueado')

    def test_produto_nao_vincula_varios_clientes(self):
        cliente_a = self._cliente()
        self._auth(self.admin)
        cliente_b = self.api.post(
            '/api/comercial/clientes/',
            {**PAYLOAD, 'cnpj': '00000000000272', 'razaoSocial': 'EMPRESA B LTDA'},
            format='json',
            **HEADERS,
        )
        self.assertEqual(cliente_b.status_code, 201, cliente_b.content)
        self._auth(self.produtos)
        created = self.api.post(
            '/api/comercial/produtos/',
            {
                'nome': 'Produto exclusivo',
                'fispq': '',
                'clienteIds': [cliente_a, cliente_b.json()['id']],
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(created.status_code, 400, created.content)
        self.assertIn('único', str(created.json()).lower())

    def test_lista_produtos_filtra_por_cliente(self):
        cliente_a = self._cliente()
        self._auth(self.admin)
        cliente_b = self.api.post(
            '/api/comercial/clientes/',
            {**PAYLOAD, 'cnpj': '00000000000353', 'razaoSocial': 'EMPRESA C LTDA'},
            format='json',
            **HEADERS,
        )
        self.assertEqual(cliente_b.status_code, 201, cliente_b.content)
        self._auth(self.produtos)
        self.api.post(
            '/api/comercial/produtos/',
            {'nome': 'Produto A', 'clienteIds': [cliente_a]},
            format='json',
            **HEADERS,
        )
        self.api.post(
            '/api/comercial/produtos/',
            {'nome': 'Produto B', 'clienteIds': [cliente_b.json()['id']]},
            format='json',
            **HEADERS,
        )
        filtrado = self.api.get(f'/api/comercial/produtos/?cliente={cliente_a}', **HEADERS)
        self.assertEqual(filtrado.status_code, 200, filtrado.content)
        nomes = [item['nome'] for item in filtrado.json()['results']]
        self.assertIn('PRODUTO A', nomes)
        self.assertNotIn('PRODUTO B', nomes)

    def _usuario_notificacao(self, username, google_email=None, **kwargs):
        defaults = {
            'password': 'test123',
            'name': username,
            'role_id': '2',
            'status': 'ativo',
            'environments': ['Comercial'],
            'funcoes': {'Comercial': ['receber-email-homologacao']},
            'google_email': google_email,
        }
        defaults.update(kwargs)
        return User.objects.create_user(username=username, **defaults)

    def _vincular_produto(self, cliente_id, nome='Glifosato 480'):
        self._auth(self.produtos)
        created = self.api.post(
            '/api/comercial/produtos/',
            {
                'nome': nome,
                'numeroOnu': '3082',
                'classeRisco': '9',
                'grupoEmbalagem': 'III',
                'fispq': 'https://exemplo.com/glifosato.pdf',
                'clienteIds': [cliente_id],
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(created.status_code, 201, created.content)
        return created

    def test_pendente_envia_email_para_quem_optou(self):
        self.produtos.google_email = 'produtos@transcamila.com.br'
        self.produtos.save(update_fields=['google_email'])
        self._usuario_notificacao('homolog.mail', 'validador@transcamila.com.br')
        cliente_id = self._cliente()
        self._vincular_produto(cliente_id)
        self.assertEqual(len(mail.outbox), 1)
        email_obj = mail.outbox[0]
        self.assertIn('digitalmidia@transcamila.com.br', email_obj.from_email)
        self.assertIn('validador@transcamila.com.br', email_obj.to)
        self.assertIn('TccConex', email_obj.subject)
        self.assertIn('GLIFOSATO', email_obj.body)
        self.assertIn('3082', email_obj.body)
        self.assertIn('Abrir para aprovar', email_obj.body)
        self.assertIn(f'validacao-clientes?cliente={cliente_id}', email_obj.body)

    def test_admin_sem_opt_in_nao_recebe_email(self):
        self.admin.google_email = 'admin@transcamila.com.br'
        self.admin.save(update_fields=['google_email'])
        self.produtos.google_email = 'produtos@transcamila.com.br'
        self.produtos.save(update_fields=['google_email'])
        cliente_id = self._cliente()
        self._vincular_produto(cliente_id)
        self.assertEqual(len(mail.outbox), 0)

    def test_ja_pendente_nao_reenvia_email(self):
        self.produtos.google_email = 'produtos@transcamila.com.br'
        self.produtos.save(update_fields=['google_email'])
        self._usuario_notificacao('homolog.mail2', 'validador2@transcamila.com.br')
        cliente_id = self._cliente()
        self._vincular_produto(cliente_id, 'Produto um')
        self.assertEqual(len(mail.outbox), 1)
        self._vincular_produto(cliente_id, 'Produto dois')
        self.assertEqual(len(mail.outbox), 1)

    def test_opt_in_sem_google_e_ignorado(self):
        self.produtos.google_email = 'produtos@transcamila.com.br'
        self.produtos.save(update_fields=['google_email'])
        self._usuario_notificacao('homolog.semgmail', google_email=None)
        self._usuario_notificacao('homolog.comgmail', 'com.gmail@transcamila.com.br')
        cliente_id = self._cliente()
        self._vincular_produto(cliente_id)
        self.assertEqual(len(mail.outbox), 1)
        email_obj = mail.outbox[0]
        self.assertEqual(email_obj.to, ['com.gmail@transcamila.com.br'])
        self.assertIn('digitalmidia@transcamila.com.br', email_obj.from_email)

    def test_lote_cadastra_varios_e_dispara_um_email(self):
        self.produtos.google_email = 'produtos@transcamila.com.br'
        self.produtos.save(update_fields=['google_email'])
        self._usuario_notificacao('homolog.lote', 'lote@transcamila.com.br')
        cliente_id = self._cliente()
        self._auth(self.produtos)
        created = self.api.post(
            '/api/comercial/produtos/lote/',
            {
                'clienteId': cliente_id,
                'produtos': [
                    {
                        'nome': 'Glifosato 480',
                        'numeroOnu': '3082',
                        'classeRisco': '9',
                        'grupoEmbalagem': 'III',
                        'fispq': 'https://exemplo.com/glifosato.pdf',
                    },
                    {
                        'nome': 'Ureia agrícola',
                        'classeRisco': 'nao_classificado',
                    },
                    {
                        'nome': 'Solvente especial',
                        'numeroOnu': '1993',
                        'classeRisco': '3',
                        'grupoEmbalagem': 'II',
                    },
                ],
            },
            format='json',
            **HEADERS,
        )
        self.assertEqual(created.status_code, 201, created.content)
        self.assertEqual(created.json()['count'], 3)
        self.assertEqual(len(mail.outbox), 1)
        email_obj = mail.outbox[0]
        self.assertIn('digitalmidia@transcamila.com.br', email_obj.from_email)
        self.assertIn('GLIFOSATO', email_obj.body)
        self.assertIn('UREIA', email_obj.body)
        self.assertIn('SOLVENTE', email_obj.body)
        self.assertIn('1993', email_obj.body)
        self._auth(self.admin)
        cliente = self.api.get(f'/api/comercial/clientes/{cliente_id}/', **HEADERS)
        self.assertEqual(cliente.json()['compatibilidade'], 'pendente_validacao')
        self.assertEqual(cliente.json()['produtosCount'], 3)

