import io
import re
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import (
    STATUS_TABELA_FRETE_CHOICES,
    TIPO_TABELA_DISTRIBUICAO,
    TIPO_TABELA_FRETE_CHOICES,
)
from .tabela_distribuicao import merge_config

_HEADER_FILL = PatternFill('solid', fgColor='118CC4')
_HEADER_FONT = Font(bold=True, color='FFFFFF')
_LABEL_FONT = Font(bold=True, color='334155')


def _as_number(value):
    if value in (None, ''):
        return None
    if isinstance(value, (int, float, Decimal)):
        return value
    texto = str(value).strip().replace('%', '').replace(' ', '')
    if not texto:
        return None
    if ',' in texto and '.' in texto:
        texto = texto.replace('.', '').replace(',', '.')
    else:
        texto = texto.replace(',', '.')
    try:
        return Decimal(texto)
    except InvalidOperation:
        return str(value)


def _fator_para_percentual(value):
    numero = _as_number(value)
    if numero is None or isinstance(numero, str):
        return numero
    return (numero * Decimal('100')).quantize(Decimal('0.0001'))


def _safe_filename(tabela):
    base = (tabela.codigo or tabela.nome_base() or 'tabela-frete').strip()
    slug = re.sub(r'[\\/:*?"<>|]+', '', base)
    slug = re.sub(r'\s+', '_', slug).strip('._') or 'tabela-frete'
    return f'{slug}_rev{tabela.revisao}.xlsx'


def _write_header_row(ws, values, row=1):
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal='center', wrap_text=True, vertical='center')
    ws.row_dimensions[row].height = 22


def _autosize(ws, min_width=12, max_width=36):
    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        width = min_width
        for cell in column:
            value = '' if cell.value is None else str(cell.value)
            width = max(width, min(max_width, len(value) + 2))
        ws.column_dimensions[letter].width = width


def _label_tipo(tabela):
    return dict(TIPO_TABELA_FRETE_CHOICES).get(tabela.tipo, tabela.tipo)


def _label_status(tabela):
    return dict(STATUS_TABELA_FRETE_CHOICES).get(tabela.status, tabela.status)


def _clientes_texto(tabela):
    nomes = [
        (cliente.razao_social or cliente.nome_fantasia or '').strip()
        for cliente in tabela.clientes.all()
    ]
    return '; '.join(nome for nome in nomes if nome) or '—'


def _sheet_identificacao(wb, tabela):
    ws = wb.active
    ws.title = 'Identificação'
    pares = [
        ('Nome', tabela.nome_base() or tabela.nome),
        ('Código', tabela.codigo or '—'),
        ('Revisão', tabela.revisao),
        ('Status', _label_status(tabela)),
        ('Tipo', _label_tipo(tabela)),
        ('Clientes', _clientes_texto(tabela)),
    ]
    _write_header_row(ws, ['Campo', 'Valor'])
    for index, (rotulo, valor) in enumerate(pares, start=2):
        label_cell = ws.cell(row=index, column=1, value=rotulo)
        label_cell.font = _LABEL_FONT
        ws.cell(row=index, column=2, value=valor)
    _autosize(ws, min_width=14, max_width=60)
    return ws


def _sheet_distribuicao(wb, tabela):
    ws = wb.create_sheet('Faixas')
    faixas = list(tabela.faixas or [])
    bandas = []
    if faixas:
        bandas = faixas[0].get('tarifas') or []
    elif isinstance(tabela.config, dict):
        bandas = tabela.config.get('bandas') or []
    headers = ['De', 'Até', 'Frete mín.']
    headers.extend(banda.get('rotulo') or banda.get('key') or f'Banda {i + 1}' for i, banda in enumerate(bandas))
    extras = []
    if faixas:
        extras = faixas[0].get('extras') or []
    elif isinstance(tabela.config, dict):
        extras = tabela.config.get('colunasExtras') or []
    config = merge_config(tabela.config or {})
    gris_adv_unificado = bool(config.get('grisAdvUnificado'))
    headers.extend(['Pedágio/t'])
    if gris_adv_unificado:
        headers.append('GRIS/ADV %')
    else:
        headers.extend(['GRIS %', 'ADV %'])
    headers.extend(extra.get('rotulo') or extra.get('key') or f'Coluna {i + 1}' for i, extra in enumerate(extras))
    headers.extend(['Prazo frac.', 'Prazo fechado'])
    _write_header_row(ws, headers)
    for row_index, faixa in enumerate(faixas, start=2):
        valores = [
            faixa.get('kmDe'),
            faixa.get('kmAte'),
            _as_number(faixa.get('freteMinimo')),
        ]
        tarifas = {item.get('key'): item.get('valor') for item in (faixa.get('tarifas') or []) if item.get('key')}
        for banda in bandas:
            valores.append(_as_number(tarifas.get(banda.get('key')) if banda.get('key') else None))
        if not bandas:
            for tarifa in faixa.get('tarifas') or []:
                valores.append(_as_number(tarifa.get('valor')))
        valores.append(_as_number(faixa.get('pedagioTon')))
        if gris_adv_unificado:
            valores.append(_fator_para_percentual(faixa.get('grisPercent')))
        else:
            valores.extend([
                _fator_para_percentual(faixa.get('grisPercent')),
                _fator_para_percentual(faixa.get('advPercent')),
            ])
        extras_faixa = {item.get('key'): item.get('valor') for item in (faixa.get('extras') or []) if item.get('key')}
        for extra in extras:
            valores.append(_as_number(extras_faixa.get(extra.get('key')) if extra.get('key') else None))
        valores.extend([
            faixa.get('prazoFracionado'),
            faixa.get('prazoFechado'),
        ])
        for col, value in enumerate(valores, start=1):
            ws.cell(row=row_index, column=col, value=value)
    _autosize(ws)
    return ws


def _sheet_transferencia(wb, tabela):
    ws = wb.create_sheet('Trechos')
    _write_header_row(ws, ['Origem', 'Destino', 'Veículo', 'Frete', 'Pedágio', 'GRIS', 'Ad-VL', 'Prazo'])
    for row_index, linha in enumerate(tabela.linhas.all(), start=2):
        valores = [
            linha.origem or '',
            linha.entrega or '',
            linha.veiculo or '',
            linha.tarifa_frete,
            linha.pedagio,
            linha.gris or '',
            linha.ad_valorem or '',
            linha.prazo_dias or '',
        ]
        for col, value in enumerate(valores, start=1):
            ws.cell(row=row_index, column=col, value=value)
    _autosize(ws)
    return ws


def build_tabela_frete_xlsx(tabela):
    wb = Workbook()
    _sheet_identificacao(wb, tabela)
    if tabela.tipo == TIPO_TABELA_DISTRIBUICAO:
        _sheet_distribuicao(wb, tabela)
    else:
        _sheet_transferencia(wb, tabela)
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue(), _safe_filename(tabela)
