"""Planilha (.xlsx) com os dados da proposta comercial, usada no repasse para a operação."""

import io
from copy import deepcopy

from openpyxl import Workbook
from openpyxl.styles import Alignment, PatternFill

from .models import TIPO_GENERALIDADE_CHOICES
from .proposta_email_service import _safe_filename_part
from .proposta_tarifas import (
    _margem_decimal,
    _margens_efetivas_auditoria,
    _margens_padrao_cliente,
    aplicar_margens_config,
    calcular_trecho,
    faixas_comerciais,
    rotulo_revisao,
    rotulos_veiculo_cliente,
    snapshot_distribuicao,
    tabela_distribuicao_do_cliente,
    variacao_margem_distribuicao,
)
from .tabela_distribuicao import (
    VEICULOS_TARIFA_MODELO_OFICIAL,
    gerar_faixas_distribuicao,
    merge_config,
    preset_config_oficial_distribuicao,
)
from .tabela_frete_export import _LABEL_FONT, _as_number, _autosize, _write_header_row

_MODALIDADE_LABEL = {
    'transferencia': 'Transferência',
    'carga_fechada': 'Carga Fechada',
    'spot': 'Spot',
    'op_portuaria': 'Logística retroportuária',
}
_FMT_MOEDA = '#,##0.00'
_FMT_PCT = '0.00%'
_CALCULO_FILL = PatternFill('solid', fgColor='0B5A8A')


def proposta_excel_filename(proposta) -> str:
    numero = _safe_filename_part(rotulo_revisao(proposta), 40) or 'proposta'
    cliente = _safe_filename_part(proposta.cliente_nome or '')
    return f'Proposta_{numero}_{cliente}.xlsx' if cliente else f'Proposta_{numero}.xlsx'


def _pct(valor):
    numero = _margem_decimal(valor)
    return float(numero) if numero is not None else None


def _preenchido(valor) -> bool:
    return valor not in (None, '') and str(valor).strip() != ''


def _margens_por_veiculo(proposta):
    """[(bandaKey, rótulo, margem padrão da tabela, margem efetiva da proposta)]."""
    if not proposta.cliente_id:
        return []
    rotulos = rotulos_veiculo_cliente(proposta.cliente_id)
    efetivas = {str(item['bandaKey']): item.get('margem') for item in _margens_efetivas_auditoria(proposta)}
    return [
        (
            str(item['bandaKey']),
            rotulos.get(item['bandaKey']) or item.get('rotulo') or item['bandaKey'],
            item.get('margem'),
            efetivas.get(str(item['bandaKey']), item.get('margem')),
        )
        for item in _margens_padrao_cliente(proposta.cliente_id)
    ]


def _calculo_trecho(proposta, linha, margens):
    modalidade = (linha.modalidade or 'transferencia').strip() or 'transferencia'
    if (
        not proposta.cliente_id
        or modalidade not in ('transferencia', 'carga_fechada', 'op_portuaria', 'spot')
        or not (linha.veiculo_key or '').strip()
        or linha.km in (None, '')
    ):
        return None
    kwargs = {
        'cliente_id': proposta.cliente_id,
        'origem': linha.origem or '',
        'destino': linha.entrega or '',
        'veiculo_key': linha.veiculo_key,
        'km': linha.km,
    }
    padrao = calcular_trecho(**kwargs, margens=None)
    aplicada = calcular_trecho(**kwargs, margens=proposta.margens_veiculo or None)
    if padrao.get('erro') or aplicada.get('erro'):
        return None
    margem = margens.get(padrao.get('veiculoKey'), (None, None))
    frete = _as_number(linha.tarifa_frete)
    tarifa_padrao = _as_number(padrao.get('tarifaFrete'))
    tarifa_calculada = _as_number(aplicada.get('tarifaFrete'))
    variacao = None
    ajuste = None
    if tarifa_padrao and frete is not None:
        variacao = float(frete) / float(tarifa_padrao) - 1
    if tarifa_calculada is not None and frete is not None:
        ajuste = 'Sim' if abs(float(frete) - float(tarifa_calculada)) >= 0.01 else 'Não'
    return {
        'faixa': padrao.get('kmFaixaAte'),
        'margem_padrao': _pct(margem[0]),
        'margem_aplicada': _pct(margem[1]),
        'tarifa_padrao': tarifa_padrao,
        'tarifa_calculada': tarifa_calculada,
        'variacao': variacao,
        'ajuste': ajuste,
    }


def _gris_adv_unificado(proposta, linhas) -> bool:
    """Mesmo critério do PDF: uma coluna GRIS/ADV quando a tabela unifica ou todos os trechos coincidem."""
    if (proposta.tabela_distribuicao or {}).get('grisAdvUnificado'):
        return True
    com_taxa = [linha for linha in linhas if _preenchido(linha.gris) or _preenchido(linha.ad_valorem)]
    return bool(com_taxa) and all((linha.gris or '').strip() == (linha.ad_valorem or '').strip() for linha in com_taxa)


def _colunas_trechos(proposta, linhas, calculos):
    """(cabeçalho, valor(linha, cálculo), formato) só com o que este tipo de proposta usa."""
    modalidades = {(linha.modalidade or 'transferencia') for linha in linhas}
    portuaria = 'op_portuaria' in modalidades
    spot = 'spot' in modalidades
    com_calculo = any(calculos)

    colunas = []
    if len(modalidades) > 1:
        colunas.append(('Operação', lambda l, c: _MODALIDADE_LABEL.get(l.modalidade or 'transferencia', l.modalidade), None))
    colunas += [
        ('Origem', lambda l, c: l.origem, None),
        ('Destino', lambda l, c: l.entrega, None),
        ('Veículo', lambda l, c: l.veiculo, None),
        ('KM', lambda l, c: _as_number(l.km), None),
        ('Frete (R$)', lambda l, c: _as_number(l.tarifa_frete), _FMT_MOEDA),
        ('Pedágio (R$)', lambda l, c: _as_number(l.pedagio), _FMT_MOEDA),
    ]
    if portuaria:
        colunas += [
            ('Retirada CTNR (R$)', lambda l, c: _as_number(l.retirada_ctnt), _FMT_MOEDA),
            ('Desova CTNR (R$)', lambda l, c: _as_number(l.desova_ctnt), _FMT_MOEDA),
        ]
    if _gris_adv_unificado(proposta, linhas):
        colunas.append(('GRIS/ADV', lambda l, c: l.gris or l.ad_valorem, None))
    else:
        colunas += [
            ('GRIS', lambda l, c: l.gris, None),
            ('Ad valorem', lambda l, c: l.ad_valorem, None),
        ]
    colunas += [
        ('ICMS', lambda l, c: l.icms, None),
        ('Prazo', lambda l, c: l.prazo_dias, None),
    ]
    if spot:
        colunas.append(('Outros valores', lambda l, c: l.outros_valores, None))
    if portuaria and any(_preenchido(linha.devolucao_container) for linha in linhas):
        colunas.append(('Devolução CTNR', lambda l, c: l.devolucao_container, None))
    if any(_preenchido(linha.observacoes) for linha in linhas):
        colunas.append(('Observações', lambda l, c: l.observacoes, None))
    colunas_proposta = len(colunas)
    if com_calculo:
        colunas += [
            ('Faixa usada (até km)', lambda l, c: c.get('faixa'), None),
            ('Margem padrão', lambda l, c: c.get('margem_padrao'), _FMT_PCT),
            ('Margem aplicada', lambda l, c: c.get('margem_aplicada'), _FMT_PCT),
            ('Tarifa tabela padrão (R$)', lambda l, c: c.get('tarifa_padrao'), _FMT_MOEDA),
            ('Tarifa calculada c/ margem (R$)', lambda l, c: c.get('tarifa_calculada'), _FMT_MOEDA),
            ('Variação s/ tabela padrão', lambda l, c: c.get('variacao'), _FMT_PCT),
            ('Ajuste manual no frete', lambda l, c: c.get('ajuste'), None),
        ]
    return colunas, colunas_proposta


def _sheet_trechos(wb, proposta, margens):
    linhas = [
        linha for linha in proposta.linhas.all().order_by('ordem', 'pk')
        if any(_preenchido(v) for v in (linha.origem, linha.entrega, linha.veiculo, linha.km, linha.tarifa_frete))
    ]
    if not linhas:
        return
    calculos = [_calculo_trecho(proposta, linha, margens) for linha in linhas]
    colunas, colunas_proposta = _colunas_trechos(proposta, linhas, calculos)
    ws = wb.create_sheet('Trechos')
    _write_header_row(ws, [cabecalho for cabecalho, _valor, _fmt in colunas])
    # Colunas de memória de cálculo (complementares) com cabeçalho em azul mais escuro.
    for col in range(colunas_proposta + 1, len(colunas) + 1):
        ws.cell(row=1, column=col).fill = _CALCULO_FILL
    for row, (linha, calculo) in enumerate(zip(linhas, calculos), start=2):
        for col, (_cabecalho, valor, fmt) in enumerate(colunas, start=1):
            conteudo = valor(linha, calculo or {})
            cell = ws.cell(row=row, column=col, value=conteudo if _preenchido(conteudo) else None)
            if fmt:
                cell.number_format = fmt
    ws.freeze_panes = 'C2'
    _autosize(ws, max_width=30)


def _faixas_distribuicao(proposta):
    faixas = (proposta.tabela_distribuicao or {}).get('faixas') or []
    if faixas or not proposta.cliente_id:
        return faixas
    # Mesmo critério do PDF: sem snapshot salvo, usa a tabela de distribuição vigente do cliente.
    snap = snapshot_distribuicao(proposta.cliente_id, proposta.margens_veiculo)
    return (snap or {}).get('faixas') or []


def _grade_calculo_transferencia(proposta):
    """Grade de faixas x veículo usada por calcular_trecho, já com as margens da proposta."""
    tabela = tabela_distribuicao_do_cliente(proposta.cliente_id)
    base = (tabela.config if tabela else None) or preset_config_oficial_distribuicao()
    nome = f'{tabela.nome} (rev. {tabela.revisao})' if tabela else 'Tabela oficial de transferência'
    config = merge_config(aplicar_margens_config(base, proposta.margens_veiculo))
    if not config.get('veiculosTarifa'):
        config['veiculosTarifa'] = deepcopy(VEICULOS_TARIFA_MODELO_OFICIAL)
    faixas = faixas_comerciais(gerar_faixas_distribuicao(config))
    rotulos = rotulos_veiculo_cliente(proposta.cliente_id)
    for faixa in faixas:
        for tarifa in faixa.get('tarifas') or []:
            if tarifa.get('key') in rotulos:
                tarifa['rotulo'] = rotulos[tarifa['key']]
    return faixas, nome


def _write_margens(ws, itens, row) -> int:
    ws.cell(row=row, column=1, value='Margens por veículo').font = _LABEL_FONT
    row += 1
    _write_header_row(ws, ['Veículo', 'Margem padrão da tabela', 'Margem da proposta', 'Alterada', 'Efeito no preço'], row=row)
    for _chave, rotulo, padrao, efetiva in itens:
        row += 1
        alterada = _margem_decimal(padrao) != _margem_decimal(efetiva)
        valores = [
            rotulo, _pct(padrao), _pct(efetiva), 'Sim' if alterada else 'Não',
            (variacao_margem_distribuicao(padrao, efetiva) or '') if alterada else 'Sem alteração',
        ]
        for col, valor in enumerate(valores, start=1):
            cell = ws.cell(row=row, column=col, value=valor)
            if col in (2, 3):
                cell.number_format = _FMT_PCT
    return row + 1


def _write_faixas(ws, faixas, row, titulo) -> None:
    ws.cell(row=row, column=1, value=titulo).font = _LABEL_FONT
    row += 1
    primeira = faixas[0]
    tarifas = primeira.get('tarifas') or []
    extras = primeira.get('extras') or []
    cabecalho = ['KM de', 'KM até', 'Frete mínimo (R$)']
    cabecalho += [f"{t.get('rotulo') or t.get('key')} (R$/{t.get('unidade') or 'ton'})" for t in tarifas]
    cabecalho += ['Pedágio (R$/ton)', 'GRIS', 'Ad valorem', 'Prazo fracionado', 'Prazo fechado']
    cabecalho += [e.get('rotulo') or e.get('key') for e in extras]
    _write_header_row(ws, cabecalho, row=row)
    col_gris = 3 + len(tarifas) + 2
    for faixa in faixas:
        row += 1
        valores = [faixa.get('kmDe'), faixa.get('kmAte'), _as_number(faixa.get('freteMinimo'))]
        valores += [_as_number(t.get('valor')) for t in faixa.get('tarifas') or []]
        valores += [
            _as_number(faixa.get('pedagioTon')), _as_number(faixa.get('grisPercent')),
            _as_number(faixa.get('advPercent')), faixa.get('prazoFracionado'), faixa.get('prazoFechado'),
        ]
        valores += [_as_number(e.get('valor')) for e in faixa.get('extras') or []]
        for col, valor in enumerate(valores, start=1):
            cell = ws.cell(row=row, column=col, value=valor)
            if col in (col_gris, col_gris + 1):
                cell.number_format = _FMT_PCT
            elif col >= 3:
                cell.number_format = _FMT_MOEDA


def _sheet_grade_calculo(wb, proposta, itens_margem):
    if not proposta.cliente_id:
        return
    if proposta.inclui_distribuicao:
        faixas = _faixas_distribuicao(proposta)
        nome = (proposta.tabela_distribuicao or {}).get('nome') or 'Tabela de distribuição'
        titulo = f'Grade de cálculo: {nome}, com as margens da proposta.'
    elif any(_calculo_trecho(proposta, linha, {}) for linha in proposta.linhas.all()):
        faixas, nome = _grade_calculo_transferencia(proposta)
        titulo = f'Grade de cálculo: {nome}, com as margens da proposta (o trecho usa a faixa "KM até").'
    else:
        return
    if not faixas:
        return
    ws = wb.create_sheet('Grade de cálculo')
    row = 1
    if itens_margem:
        row = _write_margens(ws, itens_margem, row) + 1
    _write_faixas(ws, faixas, row, titulo)
    _autosize(ws, max_width=24)


def _sheet_armazenagem(wb, proposta):
    tabela = proposta.tabela_armazenagem or {}
    if not proposta.inclui_armazenagem or not tabela:
        return
    ws = wb.create_sheet('Tabela Armazenagem')
    periodo = ' a '.join(p for p in (tabela.get('periodoInicio'), tabela.get('periodoFim')) if p) or '—'
    cabecalho = [
        ('Código', tabela.get('codigo') or '—'),
        ('Local', tabela.get('local') or '—'),
        ('Período', periodo),
        ('Unidade', tabela.get('unidade') or '—'),
    ]
    row = 1
    for campo, valor in cabecalho:
        ws.cell(row=row, column=1, value=campo).font = _LABEL_FONT
        ws.cell(row=row, column=2, value=valor)
        row += 1
    row += 1
    _write_header_row(ws, ['Tarifa', 'Valor'], row=row)
    for item in tabela.get('itens') or []:
        row += 1
        ws.cell(row=row, column=1, value=item.get('rotulo'))
        ws.cell(row=row, column=2, value=item.get('valor'))
    hora_extra = tabela.get('horaExtra') or tabela.get('hora_extra') or []
    if hora_extra:
        row += 2
        _write_header_row(ws, [tabela.get('horaExtraTitulo') or 'Hora-extra', 'Valor'], row=row)
        for item in hora_extra:
            row += 1
            ws.cell(row=row, column=1, value=item.get('periodo'))
            ws.cell(row=row, column=2, value=item.get('valor'))
    if tabela.get('expediente'):
        row += 2
        ws.cell(row=row, column=1, value=tabela['expediente'])
    ws.column_dimensions['A'].width = 70
    ws.column_dimensions['B'].width = 22


def _sheet_generalidades(wb, proposta):
    itens = [item for item in proposta.condicoes or [] if isinstance(item, dict)]
    if not itens:
        return
    ws = wb.create_sheet('Generalidades')
    tipos = dict(TIPO_GENERALIDADE_CHOICES)
    _write_header_row(ws, ['Serviço', 'Item', 'Condição'])
    for row, item in enumerate(itens, start=2):
        ws.cell(row=row, column=1, value=tipos.get(item.get('tipo'), item.get('tipo') or '—'))
        ws.cell(row=row, column=2, value=item.get('rotulo') or '')
        ws.cell(row=row, column=3, value=item.get('valor') or '').alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 36
    ws.column_dimensions['C'].width = 90


def build_proposta_excel(proposta) -> tuple[bytes, str]:
    wb = Workbook()
    wb.remove(wb.active)
    itens_margem = _margens_por_veiculo(proposta)
    margens = {chave: (padrao, efetiva) for chave, _rotulo, padrao, efetiva in itens_margem}
    _sheet_trechos(wb, proposta, margens)
    _sheet_grade_calculo(wb, proposta, itens_margem)
    _sheet_armazenagem(wb, proposta)
    _sheet_generalidades(wb, proposta)
    if not wb.sheetnames:
        _write_header_row(wb.create_sheet('Proposta'), ['Proposta', rotulo_revisao(proposta)])
    wb.active = 0
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), proposta_excel_filename(proposta)
