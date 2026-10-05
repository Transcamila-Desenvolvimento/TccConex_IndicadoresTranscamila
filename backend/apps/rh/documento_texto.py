"""Extrai texto uma vez, na inclusão. A consulta do agente não reabre o arquivo."""

from pathlib import Path

LIMITE = 50_000
MAX_PAGINAS = 30
MAX_BINARIO = 8 * 1024 * 1024

def garantir_texto(documento) -> str:
    """Lê o arquivo no máximo uma vez e grava o texto para as próximas consultas."""
    if documento.texto_extraido:
        return documento.texto or ''
    texto = ''
    if documento.arquivo:
        with documento.arquivo.open('rb') as arquivo:
            texto = extrair_texto(arquivo, documento.nome_original, documento.tamanho)
    type(documento).objects.filter(pk=documento.pk).update(texto=texto, texto_extraido=True)
    documento.texto = texto
    documento.texto_extraido = True
    return texto


def extrair_texto(arquivo, nome: str, tamanho: int = 0) -> str:
    ext = Path(nome or '').suffix.lower()
    try:
        if ext in {'.txt', '.csv'}:
            return _texto_puro(arquivo)
        if ext == '.pdf':
            return _pdf(arquivo)
        if tamanho and tamanho > MAX_BINARIO:
            return ''
        if ext == '.docx':
            return _docx(arquivo)
        if ext == '.xlsx':
            return _xlsx(arquivo)
        if ext == '.xls':
            return _xls(arquivo)
    except Exception:
        return ''
    return ''


def _cortar(partes) -> str:
    texto = ' '.join(parte for parte in partes if parte)
    texto = ' '.join(texto.split())
    return texto[:LIMITE]


def _texto_puro(arquivo) -> str:
    bruto = arquivo.read(LIMITE * 4)
    if hasattr(arquivo, 'close'):
        arquivo.close()
    return _cortar([bruto.decode('utf-8', errors='replace')])


def _pdf(arquivo) -> str:
    from pypdf import PdfReader
    reader = PdfReader(arquivo)
    partes = []
    total = 0
    for indice, pagina in enumerate(reader.pages):
        if indice >= MAX_PAGINAS or total >= LIMITE:
            break
        trecho = pagina.extract_text() or ''
        partes.append(trecho)
        total += len(trecho)
    if hasattr(arquivo, 'close'):
        arquivo.close()
    return _cortar(partes)


def _docx(arquivo) -> str:
    import docx
    document = docx.Document(arquivo)
    partes = []
    total = 0
    for paragrafo in document.paragraphs:
        if total >= LIMITE:
            break
        partes.append(paragrafo.text)
        total += len(paragrafo.text)
    if hasattr(arquivo, 'close'):
        arquivo.close()
    return _cortar(partes)


def _xlsx(arquivo) -> str:
    import openpyxl
    livro = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    partes = []
    total = 0
    try:
        for planilha in livro.worksheets:
            for linha in planilha.iter_rows(values_only=True):
                if total >= LIMITE:
                    break
                celulas = [str(valor) for valor in linha if valor not in (None, '')]
                if celulas:
                    trecho = ' '.join(celulas)
                    partes.append(trecho)
                    total += len(trecho)
    finally:
        livro.close()
    return _cortar(partes)


def _xls(arquivo) -> str:
    import xlrd
    conteudo = arquivo.read(MAX_BINARIO)
    if hasattr(arquivo, 'close'):
        arquivo.close()
    livro = xlrd.open_workbook(file_contents=conteudo)
    partes = []
    total = 0
    for planilha in livro.sheets():
        for indice in range(planilha.nrows):
            if total >= LIMITE:
                break
            celulas = [str(valor) for valor in planilha.row_values(indice) if valor not in (None, '')]
            if celulas:
                trecho = ' '.join(celulas)
                partes.append(trecho)
                total += len(trecho)
    return _cortar(partes)
